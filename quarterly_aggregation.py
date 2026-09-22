# -*- coding: utf-8 -*-
"""
====================================================================
سامانه تجمیع خودکار عملکرد ۳ ماهه (فصلی) نواحی نسرا - استان اصفهان
تجمیع چندین فایل اکسل ماهانه برای هر شهرستان (مثلاً ۳ فایل برای هر ناحیه)
مقیاس نمره‌دهی رسمی: از ۷۰ (عملکرد صفر) تا ۱۰۰ (عالی)
====================================================================
"""

import sys
import os
import re
import glob
import shutil
import subprocess

# Fix Windows console UTF-8 output
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Ensure working directory is the script folder
script_dir = os.path.dirname(os.path.abspath(__file__))
if script_dir:
    os.chdir(script_dir)

def ensure_dependencies():
    packages = {
        'openpyxl': 'openpyxl',
        'PIL': 'pillow',
        'arabic_reshaper': 'arabic-reshaper',
        'bidi': 'python-bidi'
    }
    missing = []
    for mod, pip_name in packages.items():
        try:
            __import__(mod)
        except ImportError:
            missing.append(pip_name)
            
    if missing:
        print("=" * 75)
        print(f"📦 Installing required libraries ({', '.join(missing)})...")
        print("=" * 75)
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install"] + missing + ["--quiet"])
            print("✓ Packages installed successfully.\n")
        except Exception as e:
            print(f"⚠️ Warning: Could not auto-install packages: {e}\n")

ensure_dependencies()

import openpyxl

QUARTERS = [
    (1, "Bahar (Spring)", "بهار (فروردین تا خرداد)"),
    (2, "Tabestan (Summer)", "تابستان (تیر تا شهریور)"),
    (3, "Paeiz (Fall)", "پاییز (مهر تا آذر)"),
    (4, "Zemestan (Winter)", "زمستان (دی تا اسفند)")
]

DISTRICTS = [
    'آران و بیدگل', 'امام حسین(ع)', 'امام رضا(ع)', 'امام صادق(ع)', 'امام علی(ع)',
    'اردستان', 'برخوار', 'بویین و میاندشت', 'تیران و کرون', 'جرقویه',
    'چادگان', 'خمینی شهر', 'خوانسار', 'خور و بیابانک', 'درچه',
    'دهاقان', 'سمیرم', 'شاهین شهر', 'شهرضا', 'فریدن',
    'فریدون شهر', 'فلاورجان', 'کاشان', 'کوهپایه', 'گلپایگان',
    'لنجان', 'مبارکه', 'نایین', 'نجف آباد', 'نطنز',
    'ورزنه', 'هرند'
]

DISTRICT_EN_NAMES = {
    'آران و بیدگل': 'Aran_va_Bidgol',
    'امام حسین(ع)': 'Emam_Hossein',
    'امام رضا(ع)': 'Emam_Reza',
    'امام صادق(ع)': 'Emam_Sadegh',
    'امام علی(ع)': 'Emam_Ali',
    'اردستان': 'Ardestan',
    'برخوار': 'Borkhar',
    'بویین و میاندشت': 'Boein_Miandasht',
    'تیران و کرون': 'Tiran_va_Karvan',
    'جرقویه': 'Jarghooyeh',
    'چادگان': 'Chadegan',
    'خمینی شهر': 'Khomeyni_Shahr',
    'خوانسار': 'Khansar',
    'خور و بیابانک': 'Khor_Biabanak',
    'درچه': 'Dorcheh',
    'دهاقان': 'Dehaghan',
    'سمیرم': 'Semirom',
    'شاهین شهر': 'Shahin_Shahr',
    'شهرضا': 'Shahreza',
    'فریدن': 'Fereydan',
    'فریدون شهر': 'Fereydoon_Shahr',
    'فلاورجان': 'Falavarjan',
    'کاشان': 'Kashan',
    'کوهپایه': 'Koohpayeh',
    'گلپایگان': 'Golpayegan',
    'لنجان': 'Lenjan',
    'مبارکه': 'Mobarakeh',
    'نایین': 'Naeen',
    'نجف آباد': 'Najaf_Abad',
    'نطنز': 'Natanz',
    'ورزنه': 'Varzaneh',
    'هرند': 'Harand'
}

def clean_str(s):
    if not s:
        return ""
    s = str(s).strip()
    s = s.replace('ي', 'ی').replace('ك', 'ک').replace('ة', 'ه')
    s = re.sub(r'[\(\)\[\]\{\}\.\_\-\d\s\u200c]+', '', s)
    s = s.replace('ع', '')
    return s

def clean_no_vav(s):
    return clean_str(s).replace('و', '')

def match_district_name(text):
    if not text:
        return None
    c_raw = clean_str(text)
    c_raw_nv = clean_no_vav(text)
    for d in DISTRICTS:
        cd = clean_str(d)
        if cd == c_raw or cd in c_raw:
            return d
    for d in DISTRICTS:
        cd_nv = clean_no_vav(d)
        if cd_nv == c_raw_nv or cd_nv in c_raw_nv:
            return d
    return None

def parse_number(val):
    if val is None:
        return 0
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    if not s:
        return 0
    p_digits = '۰۱۲۳۴۵۶۷۸۹'
    a_digits = '٠١٢٣٤٥٦٧٨٩'
    for i in range(10):
        s = s.replace(p_digits[i], str(i)).replace(a_digits[i], str(i))
    s = s.replace(',', '').replace('،', '')
    match = re.search(r'\d+(\.\d+)?', s)
    if match:
        try:
            return float(match.group()) if '.' in match.group() else int(match.group())
        except:
            return 0
    return 0

def extract_sheet_metrics(ws):
    if ws is None:
        return {'people_sum': 0, 'classes_count': 0, 'col_name': None}
    
    target_col = None
    target_col_name = None
    
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(1, c).value or '')
        if any(k in h for k in ['نفر', 'بازدید', 'مخاطب', 'شرکت', 'تیراژ', 'مجموع']):
            target_col = c
            target_col_name = h
            break
            
    if target_col is None:
        for c in range(1, ws.max_column + 1):
            h = str(ws.cell(1, c).value or '')
            if any(k in h for k in ['تعداد', 'صفحه', 'صفحات', 'میزان']):
                target_col = c
                target_col_name = h
                break

    total_people = 0
    active_classes = 0
    
    for r in range(2, ws.max_row + 1):
        has_act = any(ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip() != '' for c in range(2, ws.max_column + 1))
        if has_act:
            active_classes += 1
            if target_col:
                v = ws.cell(r, target_col).value
                total_people += parse_number(v)
                
    return {
        'people_sum': int(total_people),
        'classes_count': active_classes,
        'col_name': target_col_name
    }

def prompt_quarter_and_year():
    print("=" * 75)
    print("   NASRA 3-MONTH (QUARTERLY) PERFORMANCE SYSTEM - ISFAHAN")
    print("   سامانه هوشمند تجمیع عملکرد ۳ ماهه (فصلی) نواحی نسرا - استان اصفهان")
    print("=" * 75)
    print("Select Quarter (Fasl):\n")
    print("  [1] Bahar    - Spring (Farvardin, Ordibehesht, Khordad)")
    print("  [2] Tabestan - Summer (Tir, Mordad, Shahrivar)")
    print("  [3] Paeiz    - Fall   (Mehr, Aban, Azar)")
    print("  [4] Zemestan - Winter (Dey, Bahman, Esfand)")
    print("-" * 75)
    
    sel_period = QUARTERS[0][2]
    try:
        ans = input("Enter Quarter number (1-4) [Default: 1 = Bahar]: ").strip()
        if ans in ['1', '2', '3', '4']:
            sel_period = QUARTERS[int(ans) - 1][2]
    except (EOFError, KeyboardInterrupt):
        sel_period = QUARTERS[0][2]

    sel_year = "1405"
    try:
        ans_y = input("Enter Year (Sal) [Default: 1405]: ").strip()
        if ans_y:
            sel_year = ans_y
    except (EOFError, KeyboardInterrupt):
        sel_year = "1405"
        
    return sel_period, sel_year

def find_quarterly_reports_files(quarter_name):
    # Searches in reports, reports_3months, or subfolder
    candidates = [
        'reports',
        'reports_3months',
        'گزارشات_سه_ماهه',
        os.path.join('reports', quarter_name.split()[0])
    ]
    all_files = []
    found_dir = 'reports'
    for c in candidates:
        if os.path.exists(c):
            files = glob.glob(os.path.join(c, '*.xlsx'))
            files = [f for f in files if not os.path.basename(f).startswith('~$')]
            if files:
                return c, files
    os.makedirs('reports', exist_ok=True)
    return 'reports', []

def process_quarterly_reports(master_excel="تهیه کارنامه ۳ ماهه نواحی.xlsx", selected_quarter=None, selected_year=None):
    if not os.path.exists(master_excel):
        print(f"❌ Error: Master 3-month Excel file '{master_excel}' not found. Regenerating...")
        try:
            import make_quarterly_excel
        except Exception:
            pass
            
    if selected_quarter and selected_year:
        quarter = selected_quarter
        year = selected_year
    else:
        quarter, year = prompt_quarter_and_year()
        
    wb_master = openpyxl.load_workbook(master_excel)
    ws_card = wb_master['کارنامه هوشمند ۳ ماهه']
    ws_card['C3'].value = quarter
    ws_card['E3'].value = str(year)
    
    print("-" * 75)
    print(f"✓ Selected Evaluation Period: [{quarter}] - Year [{year}]")
    
    folder_name, excel_files = find_quarterly_reports_files(quarter)
    print(f"📁 Reports Folder: '{folder_name}' | Total Files Detected: {len(excel_files)}")
    print("-" * 75)
    
    if not excel_files:
        print(f"⚠️ No employee Excel reports found in '{folder_name}'.")
        print(f"💡 Place your 3 monthly files per district into '{folder_name}' and run again.")
        wb_master.save(master_excel)
        return False, quarter, year
        
    # Accumulate data per district across all files!
    extracted = {}
    
    for fpath in excel_files:
        fname = os.path.basename(fpath)
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
            detected = match_district_name(fname)
            
            ws_hozori = None
            ws_tavanmand = None
            ws_majazi = None
            ws_khalagh = None
            ws_tolid = None
            
            for sname in wb.sheetnames:
                cn = clean_str(sname)
                if 'حضوری' in cn: ws_hozori = wb[sname]
                elif 'توانمند' in cn: ws_tavanmand = wb[sname]
                elif 'مجازی' in cn or 'لایو' in cn: ws_majazi = wb[sname]
                elif 'خلاق' in cn: ws_khalagh = wb[sname]
                elif 'تولید' in cn: ws_tolid = wb[sname]
                
            if not detected:
                for ws in [ws_hozori, ws_majazi, ws_khalagh, ws_tavanmand]:
                    if ws is None: continue
                    for r in range(2, min(ws.max_row + 1, 25)):
                        for c in [3, 4, 2]:
                            m = match_district_name(ws.cell(r, c).value)
                            if m:
                                detected = m
                                break
                        if detected: break
                    if detected: break
                    
            if not detected:
                print(f"⚠️ Could not detect district for file: '{fname}'")
                continue
                
            m_hoz = extract_sheet_metrics(ws_hozori)
            m_tav = extract_sheet_metrics(ws_tavanmand)
            m_maj = extract_sheet_metrics(ws_majazi)
            m_kha = extract_sheet_metrics(ws_khalagh)
            m_tol = extract_sheet_metrics(ws_tolid)
            
            hoz_tot = m_hoz['people_sum'] + m_tav['people_sum']
            hoz_cls = m_hoz['classes_count'] + m_tav['classes_count']
            metric_hoz = hoz_tot if hoz_tot > 0 else hoz_cls
            
            metric_maj = m_maj['people_sum'] if m_maj['people_sum'] > 0 else m_maj['classes_count']
            metric_kha = m_kha['people_sum'] if m_kha['people_sum'] > 0 else m_kha['classes_count']
            metric_tol = m_tol['people_sum'] if m_tol['people_sum'] > 0 else m_tol['classes_count']
            neshast = 1 if (metric_hoz + metric_maj + metric_kha + metric_tol) > 0 else 0
            
            if detected not in extracted:
                extracted[detected] = {
                    'files_count': 0,
                    'files': [],
                    'hozori_total': 0,
                    'majazi': 0,
                    'khalagh': 0,
                    'tolid': 0,
                    'neshast': 0
                }
                
            extracted[detected]['files_count'] += 1
            extracted[detected]['files'].append(fname)
            extracted[detected]['hozori_total'] += metric_hoz
            extracted[detected]['majazi'] += metric_maj
            extracted[detected]['khalagh'] += metric_kha
            extracted[detected]['tolid'] += metric_tol
            extracted[detected]['neshast'] += neshast
            
            print(f"✓ [{detected}] (فایل {extracted[detected]['files_count']}: {fname}): +{metric_hoz} نفر حضوری | +{metric_maj} مجازی | +{metric_kha} خلاقانه")
            
        except Exception as e:
            print(f"❌ Error processing file '{fname}': {e}")

    print("-" * 75)
    print("💾 Updating 3-month cumulative metrics into Excel...")
    
    ws_rep = wb_master['گزارش عملکرد ۳ ماهه']
    row_map = {}
    for r in range(2, ws_rep.max_row + 1):
        dname = ws_rep.cell(r, 1).value
        if dname:
            row_map[clean_str(dname)] = r
            row_map[clean_no_vav(dname)] = r

    active_count = 0
    inactive_count = 0
    
    for dname in DISTRICTS:
        r = row_map.get(clean_str(dname)) or row_map.get(clean_no_vav(dname))
        if not r:
            continue
            
        if dname in extracted:
            data = extracted[dname]
            ws_rep.cell(row=r, column=2, value=data['hozori_total'])
            ws_rep.cell(row=r, column=3, value=data['majazi'])
            ws_rep.cell(row=r, column=4, value=data['khalagh'])
            ws_rep.cell(row=r, column=5, value=data['tolid'])
            ws_rep.cell(row=r, column=6, value=data['neshast'])
            ws_rep.cell(row=r, column=7, value=data['files_count'])
            active_count += 1
        else:
            # No files -> zero performance
            ws_rep.cell(row=r, column=2, value=0)
            ws_rep.cell(row=r, column=3, value=0)
            ws_rep.cell(row=r, column=4, value=0)
            ws_rep.cell(row=r, column=5, value=0)
            ws_rep.cell(row=r, column=6, value=0)
            ws_rep.cell(row=r, column=7, value=0)
            inactive_count += 1
            print(f"⭕ [{dname}]: گزارشی ارسال نشده (عملکرد ۳ ماهه ۰ منظور شد - نمره ۷۰)")

    wb_master.save(master_excel)
    
    archive_name = f"کارنامه_۳ماهه_نواحی_{quarter.split()[0]}_{year}.xlsx"
    try:
        shutil.copyfile(master_excel, archive_name)
        print(f"📁 پشتیبان فصلی با نام '{archive_name}' ذخیره شد.")
    except Exception:
        pass
        
    print(f"🎉 عملیات با موفقیت پایان یافت! آمار {active_count} ناحیه فعال و {inactive_count} ناحیه فاقد فعالیت ثبت شد.")
    return True, quarter, str(year)

def generate_all_quarterly_images(master_excel="تهیه کارنامه ۳ ماهه نواحی.xlsx", quarter="بهار (فروردین تا خرداد)", year="1405"):
    try:
        from image_generator import generate_quarterly_scorecard_png, generate_quarterly_dashboard_png
    except Exception as e:
        print("⚠️ ماژول‌های تولید تصویر در دسترس نیستند:", e)
        return
        
    quarter_slug = quarter.split()[0]
    out_dir = os.path.join("output_cards_3months", f"{quarter_slug}_{year}")
    os.makedirs(out_dir, exist_ok=True)
    
    print("-" * 75)
    print(f"📸 در حال تولید کارنامه‌های ۳ ماهه (مقیاس ۷۰ تا ۱۰۰) در پوشه '{out_dir}'...")
    
    wb = openpyxl.load_workbook(master_excel, data_only=True)
    ws_target = wb['پایگاه داده حد انتظار ۳ ماهه']
    ws_rep = wb['گزارش عملکرد ۳ ماهه']
    
    targets = {}
    for r in range(2, 34):
        dn = ws_target.cell(r, 1).value
        targets[dn] = {
            'branches': ws_target.cell(r, 2).value or 0,
            'hozori': ws_target.cell(r, 3).value or 0,
            'majazi': ws_target.cell(r, 4).value or 0,
            'khalagh': ws_target.cell(r, 5).value or 0,
            'tolid': ws_target.cell(r, 6).value or 0,
            'neshast': ws_target.cell(r, 7).value or 3
        }
        
    actuals = {}
    for r in range(2, ws_rep.max_row + 1):
        dn = ws_rep.cell(r, 1).value
        if dn:
            actuals[dn] = {
                'hozori': ws_rep.cell(r, 2).value or 0,
                'majazi': ws_rep.cell(r, 3).value or 0,
                'khalagh': ws_rep.cell(r, 4).value or 0,
                'tolid': ws_rep.cell(r, 5).value or 0,
                'neshast': ws_rep.cell(r, 6).value or 0,
                'files_count': ws_rep.cell(r, 7).value or 0
            }
            
    # Calculate Realization and 70-100 Score!
    # Formula: Score = 70.0 + (30.0 * min(1.0, max(0.0, realization_pct / 100.0)))
    dist_scores = []
    for dn, t in targets.items():
        a = actuals.get(dn, {'hozori': 0, 'majazi': 0, 'khalagh': 0, 'tolid': 0, 'neshast': 0, 'files_count': 0})
        pcts = []
        for k in ['hozori', 'majazi', 'khalagh', 'tolid', 'neshast']:
            tv = t.get(k, 1)
            av = a.get(k, 0)
            pcts.append((av / tv * 100) if tv > 0 else 0)
        avg_realization = sum(pcts) / len(pcts) if pcts else 0
        a['overall_realization'] = avg_realization
        
        # 70 to 100 Scale:
        # Zero performance -> 70.0
        # 100% performance -> 100.0
        sc_70_100 = round(70.0 + 30.0 * min(1.0, max(0.0, avg_realization / 100.0)), 1)
        a['score_70_100'] = sc_70_100
        dist_scores.append((dn, sc_70_100, avg_realization))
        
    dist_scores.sort(key=lambda x: (x[1], x[2]), reverse=True)
    
    rank_map = {}
    current_rank = 1
    for dn, sc, real in dist_scores:
        if real > 0:
            rank_map[dn] = str(current_rank)
            current_rank += 1
        else:
            rank_map[dn] = "عدم فعالیت"
            
    count_img = 0
    for dn in DISTRICTS:
        t = targets.get(dn, {})
        a = actuals.get(dn, {'overall_realization': 0, 'score_70_100': 70.0, 'files_count': 0})
        sc = a.get('score_70_100', 70.0)
        real = a.get('overall_realization', 0.0)
        rk = rank_map.get(dn, "عدم فعالیت")
        
        if sc >= 99.9:
            tier = "عالی (پیشتاز)"
        elif sc >= 92.5:
            tier = "خوب"
        elif sc >= 85.0:
            tier = "متوسط"
        elif sc > 70.0:
            tier = "ضعیف"
        else:
            tier = "فاقد عملکرد"
            
        en_name = DISTRICT_EN_NAMES.get(dn, dn)
        out_p_fa = os.path.join(out_dir, f"کارنامه_{dn}.png")
        out_p_en = os.path.join(out_dir, f"Scorecard_{en_name}.png")
        
        try:
            generate_quarterly_scorecard_png(
                district_name=dn,
                target_dict=t,
                actual_dict=a,
                rank=rk,
                tier=tier,
                period=quarter,
                year=year,
                files_count=a.get('files_count', 0),
                output_path=out_p_fa
            )
            shutil.copyfile(out_p_fa, out_p_en)
            count_img += 1
        except Exception as err:
            print(f"Error generating 3-month card for {dn}: {err}")

    # Provincial Dashboard
    sum_t_hoz = sum(targets[d]['hozori'] for d in targets)
    sum_t_maj = sum(targets[d]['majazi'] for d in targets)
    sum_t_kha = sum(targets[d]['khalagh'] for d in targets)
    sum_t_tol = sum(targets[d]['tolid'] for d in targets)
    sum_t_nes = sum(targets[d]['neshast'] for d in targets)
    
    sum_a_hoz = sum(actuals.get(d, {}).get('hozori', 0) for d in targets)
    sum_a_maj = sum(actuals.get(d, {}).get('majazi', 0) for d in targets)
    sum_a_kha = sum(actuals.get(d, {}).get('khalagh', 0) for d in targets)
    sum_a_tol = sum(actuals.get(d, {}).get('tolid', 0) for d in targets)
    sum_a_nes = sum(actuals.get(d, {}).get('neshast', 0) for d in targets)
    
    macro_data = [
        ('سواد رسانه حضوری و توانمندسازی (ضریب ۹۳)', int(sum_t_hoz), int(sum_a_hoz)),
        ('سواد رسانه مجازی و لایو (ضریب ۶۵۱)', int(sum_t_maj), int(sum_a_maj)),
        ('اقدامات و ابتکارات خلاقانه (ضریب ۱۸۶)', int(sum_t_kha), int(sum_a_kha)),
        ('تولیدات رسانه‌ای و محتوایی (ضریب ۹)', int(sum_t_tol), int(sum_a_tol)),
        ('نشست با انجمن مدرسان (۳ نشست)', int(sum_t_nes), int(sum_a_nes))
    ]
    
    top5 = [(i+1, dist_scores[i][0], dist_scores[i][1]) for i in range(min(5, len(dist_scores)))]
    bot5 = [(i+1, dist_scores[len(dist_scores)-1-i][0], dist_scores[len(dist_scores)-1-i][1]) for i in range(min(5, len(dist_scores)))]
    
    avg_score = sum(x[1] for x in dist_scores) / len(dist_scores) if dist_scores else 70.0
    top_d = dist_scores[0][0] if dist_scores and dist_scores[0][1] > 70.0 else "در انتظار"
    rep_c = sum(1 for x in dist_scores if x[1] > 70.0)
    
    kpi_d = {'avg_score': avg_score, 'top_district': top_d, 'reported_count': rep_c}
    dash_path_fa = os.path.join(out_dir, f"تصویر_داشبورد_مدیریتی_۳ماهه_{quarter_slug}_{year}.png")
    dash_path_en = os.path.join(out_dir, f"Dashboard_3Months_Provincial_{quarter_slug}.png")
    
    try:
        generate_quarterly_dashboard_png(macro_data, top5, bot5, kpi_d, period=quarter, year=year, output_path=dash_path_fa)
        shutil.copyfile(dash_path_fa, dash_path_en)
    except Exception as err:
        print(f"Error generating 3-month dashboard: {err}")
        
    print("-" * 75)
    print(f"🎉 تعداد {count_img} تصویر کارنامه ۳ ماهه (نمره ۷۰ تا ۱۰۰) در پوشه:")
    print(f"   📁 '{os.path.abspath(out_dir)}'")
    print("   به همراه تصویر داشبورد فصلی کل استان ذخیره گردید!")
    print("=" * 75)
    
    if sys.platform == 'win32':
        try:
            os.system(f'explorer "{os.path.abspath(out_dir)}"')
        except Exception:
            pass

if __name__ == '__main__':
    q_arg = None
    y_arg = None
    if len(sys.argv) > 1:
        raw_q = sys.argv[1].strip()
        if raw_q in ['1', '2', '3', '4']:
            q_arg = QUARTERS[int(raw_q) - 1][2]
        else:
            for q_num, q_en, q_fa in QUARTERS:
                if raw_q.lower() in q_en.lower() or raw_q in q_fa:
                    q_arg = q_fa
                    break
                    
    if len(sys.argv) > 2:
        y_arg = sys.argv[2].strip()
        
    success, sel_quarter, sel_year = process_quarterly_reports(selected_quarter=q_arg, selected_year=y_arg)
    generate_all_quarterly_images(quarter=sel_quarter, year=sel_year)
