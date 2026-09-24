# -*- coding: utf-8 -*-
"""
====================================================================
سامانه هوشمند سنجش نمرات دوره‌ای نواحی نسرا - استان اصفهان
(دوره‌های ۲ ماهه، ۳ ماهه و ۶ ماهه)

منطق مصوب اوزان و سقف شاخص‌ها:
- ۱۰٪ سواد رسانه حضوری
- ۳۰٪ سواد رسانه مجازی
  (قابلیت سرشکن و تبدیل حضوری و مجازی روی هم در سبد ۴۰ درصدی آموزش)
- ۵۰٪ اقدامات خلاقانه (دارای سقف قطعی ۱۰۰٪ بدون سرریز کاذب)
- ۱۰٪ تولیدات رسانه‌ای (دارای سقف قطعی ۱۰۰٪)

مقیاس نمره‌دهی: ۷۰ (عملکرد صفر) تا ۱۰۰ (تحقق کامل)
سطح کیفی ۳ گانه: عالی (۹۰-۱۰۰)، متوسط (۸۰-۸۹.۹)، ضعیف (زیر ۸۰)
====================================================================
"""

import sys
import os
import re
import glob
import shutil
import subprocess

# Fix Windows console UTF-8 output
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

script_dir = os.path.dirname(os.path.abspath(__file__))
if script_dir:
    os.chdir(script_dir)

def ensure_dependencies():
    packages = {'openpyxl': 'openpyxl'}
    missing = []
    for mod, pip_name in packages.items():
        try:
            __import__(mod)
        except ImportError:
            missing.append(pip_name)
    if missing:
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install"] + missing + ["--quiet", "--break-system-packages"])
        except Exception:
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install"] + missing + ["--quiet"])
            except Exception as e:
                print(f"⚠️ Warning: Could not auto-install packages: {e}")

ensure_dependencies()

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DISTRICT_BRANCHES = [
    ('آران و بیدگل', 10),
    ('امام حسین(ع)', 5),
    ('امام رضا(ع)', 22),
    ('امام صادق(ع)', 22),
    ('امام علی(ع)', 6),
    ('اردستان', 6),
    ('برخوار', 5),
    ('بویین و میاندشت', 3),
    ('تیران و کرون', 6),
    ('جرقویه', 4),
    ('چادگان', 5),
    ('خمینی شهر', 9),
    ('خوانسار', 3),
    ('خور و بیابانک', 2),
    ('درچه', 4),
    ('دهاقان', 3),
    ('سمیرم', 7),
    ('شاهین شهر', 9),
    ('شهرضا', 8),
    ('فریدن', 4),
    ('فریدون شهر', 3),
    ('فلاورجان', 11),
    ('کاشان', 15),
    ('کوهپایه', 4),
    ('گلپایگان', 5),
    ('لنجان', 12),
    ('مبارکه', 8),
    ('نایین', 4),
    ('نجف آباد', 14),
    ('نطنز', 4),
    ('ورزنه', 2),
    ('هرند', 4)
]

DISTRICTS = [d[0] for d in DISTRICT_BRANCHES]
BRANCH_MAP = {d[0]: d[1] for d in DISTRICT_BRANCHES}

# Monthly base multipliers per hozeh:
BASE_HOZORI = 31
BASE_MAJAZI = 217
BASE_KHALAGH = 62
BASE_TOLID = 3

# Approved Weights
WEIGHT_HOZORI = 0.10   # 10%
WEIGHT_MAJAZI = 0.30   # 30%
WEIGHT_TRAINING_POOL = 0.40 # 40% (Combined in-person & virtual)
WEIGHT_KHALAGH = 0.50  # 50% (Hard cap 100%)
WEIGHT_TOLID = 0.10    # 10% (Hard cap 100%)

PERIOD_CONFIGS = {
    2: {
        'months': 2,
        'title': '۲ ماهه',
        'title_en': '2-Month',
        'desc': 'دوره ۲ ماهه (۲ پوشه ماهانه یا ۲ فایل برای هر ناحیه)',
        'hozori_mult': BASE_HOZORI * 2,    # 62
        'majazi_mult': BASE_MAJAZI * 2,    # 434
        'khalagh_mult': BASE_KHALAGH * 2,  # 124
        'tolid_mult': BASE_TOLID * 2       # 6
    },
    3: {
        'months': 3,
        'title': '۳ ماهه',
        'title_en': '3-Month',
        'desc': 'دوره ۳ ماهه (۳ پوشه ماهانه یا ۳ فایل برای هر ناحیه)',
        'hozori_mult': BASE_HOZORI * 3,    # 93
        'majazi_mult': BASE_MAJAZI * 3,    # 651
        'khalagh_mult': BASE_KHALAGH * 3,  # 186
        'tolid_mult': BASE_TOLID * 3       # 9
    },
    6: {
        'months': 6,
        'title': '۶ ماهه',
        'title_en': '6-Month',
        'desc': 'دوره ۶ ماهه (۶ پوشه ماهانه یا ۶ فایل برای هر ناحیه)',
        'hozori_mult': BASE_HOZORI * 6,    # 186
        'majazi_mult': BASE_MAJAZI * 6,    # 1302
        'khalagh_mult': BASE_KHALAGH * 6,  # 372
        'tolid_mult': BASE_TOLID * 6       # 18
    }
}

def clean_str(s):
    if not s:
        return ""
    s = str(s).strip()
    s = s.replace('ي', 'ی').replace('ك', 'ک').replace('ة', 'ه')
    s = re.sub(r'[\(\)\[\]\{\}\.\_\-\d\s\u200c]+', '', s)
    s = s.replace('ع', '')
    return s

def clean_no_vav(s):
    return clean_str(s).replace('و', '')

def match_district_name(text):
    if not text:
        return None
    c_raw = clean_str(text)
    c_raw_nv = clean_no_vav(text)
    for d in DISTRICTS:
        cd = clean_str(d)
        if cd == c_raw or cd in c_raw:
            return d
    for d in DISTRICTS:
        cd_nv = clean_no_vav(d)
        if cd_nv == c_raw_nv or cd_nv in c_raw_nv:
            return d
    return None

def parse_number(val):
    if val is None:
        return 0
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    if not s:
        return 0
    p_digits = '۰۱۲۳۴۵۶۷۸۹'
    a_digits = '٠١٢٣٤٥٦٧٨٩'
    for i in range(10):
        s = s.replace(p_digits[i], str(i)).replace(a_digits[i], str(i))
    s = s.replace(',', '').replace('،', '')
    match = re.search(r'\d+(\.\d+)?', s)
    if match:
        try:
            return float(match.group()) if '.' in match.group() else int(match.group())
        except:
            return 0
    return 0

def extract_sheet_metrics(ws):
    if ws is None:
        return {'people_sum': 0, 'classes_count': 0}
    
    target_col = None
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(1, c).value or '')
        if any(k in h for k in ['نفر', 'بازدید', 'مخاطب', 'شرکت', 'تیراژ', 'مجموع']):
            target_col = c
            break
            
    if target_col is None:
        for c in range(1, ws.max_column + 1):
            h = str(ws.cell(1, c).value or '')
            if any(k in h for k in ['تعداد', 'صفحه', 'صفحات', 'میزان']):
                target_col = c
                break

    total_people = 0
    active_classes = 0
    for r in range(2, ws.max_row + 1):
        has_act = any(ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip() != '' for c in range(2, ws.max_column + 1))
        if has_act:
            active_classes += 1
            if target_col:
                v = ws.cell(r, target_col).value
                total_people += parse_number(v)
                
    return {
        'people_sum': int(total_people),
        'classes_count': active_classes
    }

def get_tier_3_levels(score):
    """
    ۳ سطح کیفی مصوب:
    1. عالی: نمره ۹۰ تا ۱۰۰
    2. متوسط: نمره ۸۰ تا ۸۹.۹
    3. ضعیف: نمره زیر ۸۰ (شامل نمره ۷۰ عدم فعالیت یا کسری)
    """
    if score >= 90.0:
        return "عالی"
    elif score >= 80.0:
        return "متوسط"
    else:
        return "ضعیف"

def prompt_period():
    print("=" * 80)
    print("   سامانه هوشمند سنجش نمرات دوره‌ای نواحی نسرا - استان اصفهان")
    print("   (منطق وزنی: ۱۰٪ حضوری، ۳۰٪ مجازی، ۵۰٪ خلاقانه [سقف ۱۰۰٪]، ۱۰٪ تولیدات)")
    print("=" * 80)
    print("انتخاب طول دوره ارزیابی عملکرد:\n")
    print("  [1] دوره ۲ ماهه (2-Month) - بررسی ۲ ماه (۲ پوشه ماهانه یا ۲ فایل)")
    print("  [2] دوره ۳ ماهه (3-Month) - بررسی ۳ ماه (۳ پوشه ماهانه یا ۳ فایل)")
    print("  [3] دوره ۶ ماهه (6-Month) - بررسی ۶ ماه (۶ پوشه ماهانه یا ۶ فایل)")
    print("-" * 80)
    
    choice = "2"
    try:
        user_in = input("لطفاً عدد ۱، ۲ یا ۳ را وارد نمایید [پیش‌فرض: 2]: ").strip()
        if user_in in ['1', '2', '3']:
            choice = user_in
    except (EOFError, KeyboardInterrupt):
        choice = "2"

    map_choice = {'1': 2, '2': 3, '3': 6}
    return map_choice.get(choice, 3)

def scan_reports_directory(reports_dir='reports'):
    os.makedirs(reports_dir, exist_ok=True)
    all_entries = sorted(os.listdir(reports_dir))
    subdirs = [d for d in all_entries if os.path.isdir(os.path.join(reports_dir, d)) and not d.startswith('.') and not d.startswith('__')]
    
    files_list = []
    if subdirs:
        print(f"📁 ساختار پوشه‌بندی ماهانه شناسایی شد ({len(subdirs)} پوشه در '{reports_dir}'):")
        for sdir in subdirs:
            spath = os.path.join(reports_dir, sdir)
            s_files = glob.glob(os.path.join(spath, '*.xlsx'))
            s_files = [f for f in s_files if not os.path.basename(f).startswith('~$')]
            print(f"   📂 پوشه ماهانه «{sdir}»: شامل {len(s_files)} فایل اکسل")
            for f in s_files:
                files_list.append((f, sdir, os.path.basename(f)))
    else:
        flat_files = glob.glob(os.path.join(reports_dir, '*.xlsx'))
        flat_files = [f for f in flat_files if not os.path.basename(f).startswith('~$')]
        if flat_files:
            print(f"📁 ساختار فایل‌های مستقیم در پوشه '{reports_dir}' ({len(flat_files)} فایل اکسل)")
            for f in flat_files:
                files_list.append((f, '', os.path.basename(f)))
        else:
            print(f"📁 پوشه '{reports_dir}' آماده است (فایلی در پوشه قرار ندارد).")
            
    return subdirs, files_list

def run_period_evaluation(selected_months=None):
    if selected_months in [2, 3, 6]:
        n_months = selected_months
    else:
        n_months = prompt_period()

    cfg = PERIOD_CONFIGS[n_months]
    print("\n" + "=" * 80)
    print(f"📌 دوره انتخابی: عملکرد {cfg['title']} ({cfg['title_en']})")
    print(f"🎯 حدانتظارها: ضریب {n_months} برابری اهداف ماهانه")
    print("⚖️ اوزان ارزیابی: ۱۰٪ حضوری | ۳۰٪ مجازی (سرشکن در سبد ۴۰٪ آموزش) | ۵۰٪ خلاقانه (سقف ۱۰۰٪) | ۱۰٪ تولیدات")
    print("=" * 80)

    subdirs, all_files = scan_reports_directory('reports')
    print("-" * 80)

    accumulated = {}
    for dn in DISTRICTS:
        accumulated[dn] = {
            'files_count': 0,
            'months_found': [],
            'files': [],
            'hozori': 0,
            'majazi': 0,
            'khalagh': 0,
            'tolid': 0
        }

    for fpath, folder_label, fname in all_files:
        try:
            detected = match_district_name(fname)
            wb = openpyxl.load_workbook(fpath, data_only=True)
            
            ws_hoz = None
            ws_tav = None
            ws_maj = None
            ws_kha = None
            ws_tol = None

            for sname in wb.sheetnames:
                cn = clean_str(sname)
                if 'حضوری' in cn: ws_hoz = wb[sname]
                elif 'توانمند' in cn: ws_tav = wb[sname]
                elif 'مجازی' in cn or 'لایو' in cn: ws_maj = wb[sname]
                elif 'خلاق' in cn: ws_kha = wb[sname]
                elif 'تولید' in cn: ws_tol = wb[sname]

            if not detected:
                for ws in [ws_hoz, ws_maj, ws_kha, ws_tav]:
                    if ws is None: continue
                    for r in range(2, min(ws.max_row + 1, 25)):
                        for c in [3, 4, 2]:
                            m = match_district_name(ws.cell(r, c).value)
                            if m:
                                detected = m
                                break
                        if detected: break
                    if detected: break

            if not detected:
                print(f"⚠️ شناسایی ناحیه برای فایل '{fname}' (در پوشه '{folder_label}') ناموفق بود.")
                continue

            m_hoz = extract_sheet_metrics(ws_hoz)
            m_tav = extract_sheet_metrics(ws_tav)
            m_maj = extract_sheet_metrics(ws_maj)
            m_kha = extract_sheet_metrics(ws_kha)
            m_tol = extract_sheet_metrics(ws_tol)

            hoz_val = (m_hoz['people_sum'] + m_tav['people_sum'])
            if hoz_val == 0:
                hoz_val = m_hoz['classes_count'] + m_tav['classes_count']

            maj_val = m_maj['people_sum'] if m_maj['people_sum'] > 0 else m_maj['classes_count']
            kha_val = m_kha['people_sum'] if m_kha['people_sum'] > 0 else m_kha['classes_count']
            tol_val = m_tol['people_sum'] if m_tol['people_sum'] > 0 else m_tol['classes_count']

            month_tag = folder_label if folder_label else f"فایل {accumulated[detected]['files_count'] + 1}"
            accumulated[detected]['files_count'] += 1
            if month_tag not in accumulated[detected]['months_found']:
                accumulated[detected]['months_found'].append(month_tag)
            accumulated[detected]['files'].append(f"{folder_label}/{fname}" if folder_label else fname)
            accumulated[detected]['hozori'] += hoz_val
            accumulated[detected]['majazi'] += maj_val
            accumulated[detected]['khalagh'] += kha_val
            accumulated[detected]['tolid'] += tol_val

            src_info = f"پوشه «{folder_label}»" if folder_label else fname
            print(f"✓ [{detected}] ({src_info}): +{hoz_val} حضوری | +{maj_val} مجازی | +{kha_val} خلاقانه | +{tol_val} تولید")

        except Exception as e:
            print(f"❌ خطا در پردازش فایل '{fname}': {e}")

    print("-" * 80)
    print("📊 محاسبه نمرات بر مبنای اوزان مصوب و سقف ۱۰۰٪...")

    results = []
    for dn in DISTRICTS:
        b_count = BRANCH_MAP[dn]
        t_hoz = b_count * cfg['hozori_mult']
        t_maj = b_count * cfg['majazi_mult']
        t_kha = b_count * cfg['khalagh_mult']
        t_tol = b_count * cfg['tolid_mult']

        acc = accumulated[dn]
        a_hoz = acc['hozori']
        a_maj = acc['majazi']
        a_kha = acc['khalagh']
        a_tol = acc['tolid']
        f_cnt = acc['files_count']
        m_found = acc['months_found']

        missing_months = [s for s in subdirs if s not in m_found] if subdirs else []

        # 1. Raw realization rates
        r_hoz_raw = (a_hoz / t_hoz) if t_hoz > 0 else 0.0
        r_maj_raw = (a_maj / t_maj) if t_maj > 0 else 0.0
        r_kha_raw = (a_kha / t_kha) if t_kha > 0 else 0.0
        r_tol_raw = (a_tol / t_tol) if t_tol > 0 else 0.0

        pct_hoz_raw = r_hoz_raw * 100.0
        pct_maj_raw = r_maj_raw * 100.0
        pct_kha_raw = r_kha_raw * 100.0
        pct_tol_raw = r_tol_raw * 100.0

        # 2. Training Basket (40% Total): 10% Hozori + 30% Majazi
        # In-person and virtual offset each other smoothly up to the 40% training cap
        raw_training_share = (WEIGHT_HOZORI * r_hoz_raw) + (WEIGHT_MAJAZI * r_maj_raw)
        training_share = min(WEIGHT_TRAINING_POOL, raw_training_share)
        surplus_training = max(0.0, raw_training_share - WEIGHT_TRAINING_POOL)

        # 3. Creative (50% Weight):
        # Own creative realization + surplus from training, strictly capped at 50% (100% of creative)
        raw_khalagh_share = WEIGHT_KHALAGH * r_kha_raw
        khalagh_share = min(WEIGHT_KHALAGH, raw_khalagh_share + surplus_training)

        # 4. Productions (10% Weight): Capped strictly at 10%
        tolid_share = WEIGHT_TOLID * min(1.0, r_tol_raw)

        # Total Realization (0.0 to 1.0)
        total_realization_ratio = training_share + khalagh_share + tolid_share
        total_realization_pct = round(total_realization_ratio * 100.0, 2)

        # Final Score in 70.0 to 100.0 scale:
        final_score = round(70.0 + (30.0 * total_realization_ratio), 1)
        tier = get_tier_3_levels(final_score)

        if f_cnt >= n_months:
            status_desc = f"کامل ({len(m_found)} از {n_months} ماه)"
        elif f_cnt > 0:
            if missing_months:
                status_desc = f"کسری: {len(m_found)} از {n_months} ماه (عدم فعالیت در: {'، '.join(missing_months)})"
            else:
                status_desc = f"کسری: {f_cnt} از {n_months} ماه تحویل شده"
        else:
            status_desc = f"فاقد گزارش (عملکرد ۰ در کل {n_months} ماه)"

        results.append({
            'district': dn,
            'branches': b_count,
            'files_count': f_cnt,
            'months_found': m_found,
            'missing_months': missing_months,
            'status_desc': status_desc,
            'score': final_score,
            'tier': tier,
            'total_realization_pct': total_realization_pct,
            'training_share_pct': round(training_share * 100.0, 2),
            'surplus_training_pct': round(surplus_training * 100.0, 2),
            'khalagh_share_pct': round(khalagh_share * 100.0, 2),
            'tolid_share_pct': round(tolid_share * 100.0, 2),
            't_hoz': t_hoz, 'a_hoz': a_hoz, 'pct_hoz_raw': pct_hoz_raw,
            't_maj': t_maj, 'a_maj': a_maj, 'pct_maj_raw': pct_maj_raw,
            't_kha': t_kha, 'a_kha': a_kha, 'pct_kha_raw': pct_kha_raw,
            't_tol': t_tol, 'a_tol': a_tol, 'pct_tol_raw': pct_tol_raw
        })

    sorted_by_score = sorted(results, key=lambda x: (x['score'], x['total_realization_pct']), reverse=True)
    rank_map = {}
    active_rank = 1
    for item in sorted_by_score:
        dn = item['district']
        if item['score'] > 70.0:
            rank_map[dn] = active_rank
            active_rank += 1
        else:
            rank_map[dn] = "-"

    for item in results:
        item['rank'] = rank_map[item['district']]

    excel_filename = "نمرات_نهایی_نواحی.xlsx"
    backup_period_file = f"نمرات_عملکرد_{n_months}ماهه.xlsx"

    wb = openpyxl.Workbook()

    font_title = Font(name='Calibri', size=15, bold=True, color='1E3A8A')
    font_th = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    font_td = Font(name='Calibri', size=11, bold=False, color='0F172A')
    font_score = Font(name='Calibri', size=12, bold=True, color='047857')
    font_copy_th = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    font_copy_dn = Font(name='Calibri', size=12, bold=True, color='0F172A')
    font_copy_sc = Font(name='Calibri', size=12, bold=True, color='047857')

    fill_th_navy = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    fill_copy_header = PatternFill(start_color='059669', end_color='059669', fill_type='solid')
    
    fill_tier_ali = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')      # Green
    fill_tier_motevaset = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')# Yellow
    fill_tier_zaeef = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')     # Red

    fill_score_col = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')

    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    # ==========================================
    # SHEET 1: فقط نام و نمره (ساده‌ترین حالت کپی مستقیم)
    # ==========================================
    ws_raw = wb.active
    ws_raw.title = "فقط نام و نمره (ساده)"
    ws_raw.views.sheetView[0].rightToLeft = True
    
    ws_raw['A1'] = "نام ناحیه"
    ws_raw['B1'] = "نمره (۷۰-۱۰۰)"
    ws_raw['A1'].font = font_copy_th
    ws_raw['B1'].font = font_copy_th
    ws_raw['A1'].fill = fill_copy_header
    ws_raw['B1'].fill = fill_copy_header
    ws_raw['A1'].alignment = align_center
    ws_raw['B1'].alignment = align_center
    ws_raw['A1'].border = border_thin
    ws_raw['B1'].border = border_thin

    for idx, r in enumerate(results, start=2):
        ws_raw.cell(row=idx, column=1, value=r['district']).alignment = align_right
        ws_raw.cell(row=idx, column=2, value=r['score']).alignment = align_center
        ws_raw.cell(row=idx, column=1).font = font_copy_dn
        ws_raw.cell(row=idx, column=2).font = font_copy_sc
        ws_raw.cell(row=idx, column=1).border = border_thin
        ws_raw.cell(row=idx, column=2).border = border_thin

    ws_raw.column_dimensions['A'].width = 25
    ws_raw.column_dimensions['B'].width = 18

    # ==========================================
    # SHEET 2: جدول نمرات و تحلیل ماه‌ها
    # ==========================================
    ws_copy = wb.create_sheet(title="جدول نمرات و تحلیل ماه‌ها")
    ws_copy.views.sheetView[0].rightToLeft = True

    ws_copy.merge_cells('A1:G1')
    ws_copy['A1'] = f"جدول ارزیابی عملکرد {cfg['title']} نواحی نسرا (۱۰٪ حضوری، ۳۰٪ مجازی، ۵۰٪ خلاقانه، ۱۰٪ تولیدات)"
    ws_copy['A1'].font = font_title
    ws_copy['A1'].alignment = align_center

    headers_s2 = [
        ('ردیف', 8), ('نام ناحیه (شهرستان)', 24), ('نمره عملکرد (۷۰-۱۰۰)', 22),
        ('سطح کیفی (۳ سطح)', 18), ('رتبه استانی', 14), ('تعداد ماه‌های ارسالی', 20),
        ('وضعیت و ماه‌های کارنکرده', 40)
    ]

    for c_idx, (h_title, w) in enumerate(headers_s2, start=1):
        cell = ws_copy.cell(row=3, column=c_idx, value=h_title)
        cell.font = font_copy_th
        cell.fill = fill_copy_header
        cell.alignment = align_center
        cell.border = border_thin
        ws_copy.column_dimensions[get_column_letter(c_idx)].width = w

    for idx, r in enumerate(results, start=1):
        row_num = 3 + idx
        ws_copy.cell(row=row_num, column=1, value=idx).alignment = align_center
        ws_copy.cell(row=row_num, column=2, value=r['district']).alignment = align_right
        ws_copy.cell(row=row_num, column=3, value=r['score']).alignment = align_center
        ws_copy.cell(row=row_num, column=4, value=r['tier']).alignment = align_center
        ws_copy.cell(row=row_num, column=5, value=r['rank']).alignment = align_center
        ws_copy.cell(row=row_num, column=6, value=f"{r['files_count']} از {n_months} ماه").alignment = align_center
        ws_copy.cell(row=row_num, column=7, value=r['status_desc']).alignment = align_right

        ws_copy.cell(row=row_num, column=1).font = font_td
        ws_copy.cell(row=row_num, column=2).font = font_copy_dn
        ws_copy.cell(row=row_num, column=3).font = font_copy_sc
        ws_copy.cell(row=row_num, column=3).fill = fill_score_col

        tier_cell = ws_copy.cell(row=row_num, column=4)
        if r['tier'] == 'عالی': tier_cell.fill = fill_tier_ali
        elif r['tier'] == 'متوسط': tier_cell.fill = fill_tier_motevaset
        else: tier_cell.fill = fill_tier_zaeef

        ws_copy.cell(row=row_num, column=5).font = font_td
        ws_copy.cell(row=row_num, column=6).font = font_td
        ws_copy.cell(row=row_num, column=7).font = font_td

        for c in range(1, 8):
            ws_copy.cell(row=row_num, column=c).border = border_thin

    # ==========================================
    # SHEET 3: جدول رتبه‌بندی استانی
    # ==========================================
    ws_rank = wb.create_sheet(title="رتبه‌بندی استانی")
    ws_rank.views.sheetView[0].rightToLeft = True

    ws_rank.merge_cells('A1:F1')
    ws_rank['A1'] = f"رتبه‌بندی استانی عملکرد {cfg['title']} ۳۲ شهرستان (به ترتیب رتبه)"
    ws_rank['A1'].font = font_title
    ws_rank['A1'].alignment = align_center

    headers_s3 = [
        ('رتبه', 10), ('نام ناحیه (شهرستان)', 24), ('نمره عملکرد (۷۰-۱۰۰)', 22),
        ('سطح کیفی', 16), ('درصد تحقق وزنی', 18), ('تعداد ماه‌های ارسالی', 20)
    ]

    for c_idx, (h_title, w) in enumerate(headers_s3, start=1):
        cell = ws_rank.cell(row=3, column=c_idx, value=h_title)
        cell.font = font_th
        cell.fill = fill_th_navy
        cell.alignment = align_center
        cell.border = border_thin
        ws_rank.column_dimensions[get_column_letter(c_idx)].width = w

    for idx, r in enumerate(sorted_by_score, start=1):
        row_num = 3 + idx
        ws_rank.cell(row=row_num, column=1, value=r['rank']).alignment = align_center
        ws_rank.cell(row=row_num, column=2, value=r['district']).alignment = align_right
        ws_rank.cell(row=row_num, column=3, value=r['score']).alignment = align_center
        ws_rank.cell(row=row_num, column=4, value=r['tier']).alignment = align_center
        ws_rank.cell(row=row_num, column=5, value=f"{r['total_realization_pct']:.1f}%").alignment = align_center
        ws_rank.cell(row=row_num, column=6, value=f"{r['files_count']} از {n_months} ماه").alignment = align_center

        ws_rank.cell(row=row_num, column=1).font = font_td
        ws_rank.cell(row=row_num, column=2).font = font_copy_dn
        ws_rank.cell(row=row_num, column=3).font = font_copy_sc
        ws_rank.cell(row=row_num, column=3).fill = fill_score_col

        tier_cell = ws_rank.cell(row=row_num, column=4)
        if r['tier'] == 'عالی': tier_cell.fill = fill_tier_ali
        elif r['tier'] == 'متوسط': tier_cell.fill = fill_tier_motevaset
        else: tier_cell.fill = fill_tier_zaeef

        ws_rank.cell(row=row_num, column=5).font = font_td
        ws_rank.cell(row=row_num, column=6).font = font_td

        for c in range(1, 7):
            ws_rank.cell(row=row_num, column=c).border = border_thin

    # ==========================================
    # SHEET 4: ریز مستندات اوزان و شاخص‌ها
    # ==========================================
    ws_full = wb.create_sheet(title="ریز مستندات و اوزان شاخص‌ها")
    ws_full.views.sheetView[0].rightToLeft = True

    ws_full.merge_cells('A1:T1')
    ws_full['A1'] = f"ریز مستندات و اوزان شاخص‌های دوره {cfg['title']} به تفکیک شهرستان‌ها"
    ws_full['A1'].font = font_title
    ws_full['A1'].alignment = align_center

    headers_full = [
        ('ردیف', 6), ('نام ناحیه', 20), ('نمره (۷۰-۱۰۰)', 14), ('سطح', 12), ('رتبه', 8),
        ('ماه‌ها', 10), ('حوزه', 8),
        ('انتظار حضوری', 14), ('عملکرد حضوری', 14), ('تحقق حضوری', 12),
        ('انتظار مجازی', 14), ('عملکرد مجازی', 14), ('تحقق مجازی', 12),
        ('سهم آموزش (۴۰٪)', 14),
        ('انتظار خلاقانه', 14), ('عملکرد خلاقانه', 14), ('سهم خلاقانه (۵۰٪)', 16),
        ('انتظار تولید', 14), ('عملکرد تولید', 14), ('سهم تولید (۱۰٪)', 14)
    ]

    for c_idx, (h_title, w) in enumerate(headers_full, start=1):
        cell = ws_full.cell(row=3, column=c_idx, value=h_title)
        cell.font = font_th
        cell.fill = fill_th_navy
        cell.alignment = align_center
        cell.border = border_thin
        ws_full.column_dimensions[get_column_letter(c_idx)].width = w

    for idx, r in enumerate(results, start=1):
        row_num = 3 + idx
        ws_full.cell(row=row_num, column=1, value=idx)
        ws_full.cell(row=row_num, column=2, value=r['district'])
        ws_full.cell(row=row_num, column=3, value=r['score'])
        ws_full.cell(row=row_num, column=4, value=r['tier'])
        ws_full.cell(row=row_num, column=5, value=r['rank'])
        ws_full.cell(row=row_num, column=6, value=f"{r['files_count']} از {n_months}")
        ws_full.cell(row=row_num, column=7, value=r['branches'])
        ws_full.cell(row=row_num, column=8, value=r['t_hoz'])
        ws_full.cell(row=row_num, column=9, value=r['a_hoz'])
        ws_full.cell(row=row_num, column=10, value=f"{r['pct_hoz_raw']:.1f}%")
        ws_full.cell(row=row_num, column=11, value=r['t_maj'])
        ws_full.cell(row=row_num, column=12, value=r['a_maj'])
        ws_full.cell(row=row_num, column=13, value=f"{r['pct_maj_raw']:.1f}%")
        ws_full.cell(row=row_num, column=14, value=f"{r['training_share_pct']:.2f}%")
        ws_full.cell(row=row_num, column=15, value=r['t_kha'])
        ws_full.cell(row=row_num, column=16, value=r['a_kha'])
        ws_full.cell(row=row_num, column=17, value=f"{r['khalagh_share_pct']:.2f}%")
        ws_full.cell(row=row_num, column=18, value=r['t_tol'])
        ws_full.cell(row=row_num, column=19, value=r['a_tol'])
        ws_full.cell(row=row_num, column=20, value=f"{r['tolid_share_pct']:.2f}%")

        for c in range(1, 21):
            cell = ws_full.cell(row=row_num, column=c)
            cell.font = font_td
            cell.border = border_thin
            cell.alignment = align_right if c == 2 else align_center

        ws_full.cell(row=row_num, column=3).font = font_score
        ws_full.cell(row=row_num, column=3).fill = fill_score_col

    wb.save(excel_filename)
    try:
        shutil.copyfile(excel_filename, backup_period_file)
    except Exception:
        pass

    # Print Clean Console Output
    print("\n" + "=" * 95)
    print(f"📋 جدول نمرات دوره {cfg['title']} نواحی نسرا (اوزان: ۱۰٪ حضوری، ۳۰٪ مجازی، ۵۰٪ خلاقانه [سقف ۱۰۰٪]، ۱۰٪ تولیدات):")
    print("=" * 95)
    print(f"{'ردیف':^6} | {'نام ناحیه (شهرستان)':<20} | {'نمره':^8} | {'سطح کیفی':^10} | {'رتبه':^6} | {'تحقق کل':^10} | {'وضعیت ماه‌های ارسالی':<30}")
    print("-" * 95)
    for idx, r in enumerate(results, start=1):
        print(f"{idx:^6} | {r['district']:<20} | {r['score']:^8.1f} | {r['tier']:^10} | {str(r['rank']):^6} | {r['total_realization_pct']:^8.1f}% | {r['status_desc']:<30}")
    print("=" * 95)

    print(f"\n🎉 فایل اکسل متمرکز با موفقیت تولید شد:")
    print(f"   📄 «{os.path.abspath(excel_filename)}»")
    print(f"   (یک کپی با نام «{backup_period_file}» نیز ذخیره شد)")
    print(f"\n💡 در شیت ۱ («فقط نام و نمره»)، ستون‌ها آماده انتخاب و کپی (Ctrl+C) هستند.")
    print("=" * 95)

    if sys.platform == 'win32':
        try:
            os.system(f'start excel "{os.path.abspath(excel_filename)}"')
        except Exception:
            try:
                os.system(f'explorer "{os.path.abspath(excel_filename)}"')
            except Exception:
                pass

if __name__ == '__main__':
    arg_m = None
    if len(sys.argv) > 1:
        raw_m = sys.argv[1].strip()
        if raw_m in ['2', '3', '6']:
            arg_m = int(raw_m)
    run_period_evaluation(selected_months=arg_m)
