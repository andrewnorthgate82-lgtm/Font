# -*- coding: utf-8 -*-
"""
====================================================================
سامانه هوشمند تجمیع خودکار فایل‌های گزارش ماهانه کارمندان نواحی نسرا
استان اصفهان - سال ۱۴۰۵
پشتیبانی از تفکیک کامل ماه و سال، پوشه‌بندی اختصاصی هر ماه،
ثبت صفر برای نواحی بدون فعالیت و صدور تصاویر عمودی ۱۰۸۰×۱۹۲۰ موبایل
====================================================================
"""

import sys
import os
import re
import glob
import shutil
import subprocess

# Fix Windows console UTF-8 output
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Ensure working directory is the script folder
script_dir = os.path.dirname(os.path.abspath(__file__))
if script_dir:
    os.chdir(script_dir)

def ensure_dependencies():
    packages = {
        'openpyxl': 'openpyxl',
        'PIL': 'pillow',
        'arabic_reshaper': 'arabic-reshaper',
        'bidi': 'python-bidi'
    }
    missing = []
    for mod, pip_name in packages.items():
        try:
            __import__(mod)
        except ImportError:
            missing.append(pip_name)
            
    if missing:
        print("=" * 75)
        print(f"📦 Installing required libraries ({', '.join(missing)})...")
        print("=" * 75)
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install"] + missing + ["--quiet"])
            print("✓ Packages installed successfully.\n")
        except Exception as e:
            print(f"⚠️ Warning: Could not auto-install packages: {e}\n")

ensure_dependencies()

import openpyxl

MONTH_INFO = [
    (1, "Farvardin", "فروردین"),
    (2, "Ordibehesht", "اردیبهشت"),
    (3, "Khordad", "خرداد"),
    (4, "Tir", "تیر"),
    (5, "Mordad", "مرداد"),
    (6, "Shahrivar", "شهریور"),
    (7, "Mehr", "مهر"),
    (8, "Aban", "آبان"),
    (9, "Azar", "آذر"),
    (10, "Dey", "دی"),
    (11, "Bahman", "بهمن"),
    (12, "Esfand", "اسفند")
]

MONTH_MAP = {}
for num, en_name, fa_name in MONTH_INFO:
    MONTH_MAP[str(num)] = fa_name
    MONTH_MAP[en_name.lower()] = fa_name
    MONTH_MAP[fa_name] = fa_name

DISTRICTS = [
    'آران و بیدگل', 'امام حسین(ع)', 'امام رضا(ع)', 'امام صادق(ع)', 'امام علی(ع)',
    'اردستان', 'برخوار', 'بویین و میاندشت', 'تیران و کرون', 'جرقویه',
    'چادگان', 'خمینی شهر', 'خوانسار', 'خور و بیابانک', 'درچه',
    'دهاقان', 'سمیرم', 'شاهین شهر', 'شهرضا', 'فریدن',
    'فریدون شهر', 'فلاورجان', 'کاشان', 'کوهپایه', 'گلپایگان',
    'لنجان', 'مبارکه', 'نایین', 'نجف آباد', 'نطنز',
    'ورزنه', 'هرند'
]

DISTRICT_EN_NAMES = {
    'آران و بیدگل': 'Aran_va_Bidgol',
    'امام حسین(ع)': 'Emam_Hossein',
    'امام رضا(ع)': 'Emam_Reza',
    'امام صادق(ع)': 'Emam_Sadegh',
    'امام علی(ع)': 'Emam_Ali',
    'اردستان': 'Ardestan',
    'برخوار': 'Borkhar',
    'بویین و میاندشت': 'Boein_Miandasht',
    'تیران و کرون': 'Tiran_va_Karvan',
    'جرقویه': 'Jarghooyeh',
    'چادگان': 'Chadegan',
    'خمینی شهر': 'Khomeyni_Shahr',
    'خوانسار': 'Khansar',
    'خور و بیابانک': 'Khor_Biabanak',
    'درچه': 'Dorcheh',
    'دهاقان': 'Dehaghan',
    'سمیرم': 'Semirom',
    'شاهین شهر': 'Shahin_Shahr',
    'شهرضا': 'Shahreza',
    'فریدن': 'Fereydan',
    'فریدون شهر': 'Fereydoon_Shahr',
    'فلاورجان': 'Falavarjan',
    'کاشان': 'Kashan',
    'کوهپایه': 'Koohpayeh',
    'گلپایگان': 'Golpayegan',
    'لنجان': 'Lenjan',
    'مبارکه': 'Mobarakeh',
    'نایین': 'Naeen',
    'نجف آباد': 'Najaf_Abad',
    'نطنز': 'Natanz',
    'ورزنه': 'Varzaneh',
    'هرند': 'Harand'
}

def clean_str(s):
    if not s:
        return ""
    s = str(s).strip()
    s = s.replace('ي', 'ی').replace('ك', 'ک').replace('ة', 'ه')
    s = re.sub(r'[\(\)\[\]\{\}\.\_\-\d\s\u200c]+', '', s)
    s = s.replace('ع', '')
    return s

def clean_no_vav(s):
    return clean_str(s).replace('و', '')

def match_district_name(text):
    if not text:
        return None
    c_raw = clean_str(text)
    c_raw_nv = clean_no_vav(text)
    for d in DISTRICTS:
        cd = clean_str(d)
        if cd == c_raw or cd in c_raw:
            return d
    for d in DISTRICTS:
        cd_nv = clean_no_vav(d)
        if cd_nv == c_raw_nv or cd_nv in c_raw_nv:
            return d
    return None

def parse_number(val):
    if val is None:
        return 0
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    if not s:
        return 0
    p_digits = '۰۱۲۳۴۵۶۷۸۹'
    a_digits = '٠١٢٣٤٥٦٧٨٩'
    for i in range(10):
        s = s.replace(p_digits[i], str(i)).replace(a_digits[i], str(i))
    s = s.replace(',', '').replace('،', '')
    match = re.search(r'\d+(\.\d+)?', s)
    if match:
        try:
            return float(match.group()) if '.' in match.group() else int(match.group())
        except:
            return 0
    return 0

def extract_sheet_metrics(ws):
    if ws is None:
        return {'people_sum': 0, 'classes_count': 0, 'col_name': None}
    
    target_col = None
    target_col_name = None
    
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(1, c).value or '')
        if any(k in h for k in ['نفر', 'بازدید', 'مخاطب', 'شرکت', 'تیراژ', 'مجموع']):
            target_col = c
            target_col_name = h
            break
            
    if target_col is None:
        for c in range(1, ws.max_column + 1):
            h = str(ws.cell(1, c).value or '')
            if any(k in h for k in ['تعداد', 'صفحه', 'صفحات', 'میزان']):
                target_col = c
                target_col_name = h
                break

    total_people = 0
    active_classes = 0
    
    for r in range(2, ws.max_row + 1):
        has_act = any(ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip() != '' for c in range(2, ws.max_column + 1))
        if has_act:
            active_classes += 1
            if target_col:
                v = ws.cell(r, target_col).value
                total_people += parse_number(v)
                
    return {
        'people_sum': int(total_people),
        'classes_count': active_classes,
        'col_name': target_col_name
    }

def prompt_month_and_year(default_month='شهریور', default_year='1405'):
    print("=" * 75)
    print("   NASRA DISTRICT PERFORMANCE MONITORING & SCORECARD SYSTEM")
    print("   سامانه هوشمند تجمیع عملکرد و صدور کارنامه نواحی نسرا - استان اصفهان")
    print("=" * 75)
    print("\nSelect Evaluation Month (Shomare Mah ra entekhab konid):\n")
    print("  [1] Farvardin   (فروردین)      [7]  Mehr        (مهر)")
    print("  [2] Ordibehesht (اردیبهشت)     [8]  Aban        (آبان)")
    print("  [3] Khordad     (خرداد)        [9]  Azar        (آذر)")
    print("  [4] Tir         (تیر)          [10] Dey         (دی)")
    print("  [5] Mordad      (مرداد)        [11] Bahman      (بهمن)")
    print("  [6] Shahrivar   (شهریور)       [12] Esfand      (اسفند)")
    print("-" * 75)
    
    selected_m = default_month
    try:
        user_choice = input(f"Enter month number (e.g. 4 for Tir, 5 for Mordad) [Default: 6 = Shahrivar]: ").strip()
        if user_choice:
            user_choice_lower = user_choice.lower()
            if user_choice_lower in MONTH_MAP:
                selected_m = MONTH_MAP[user_choice_lower]
            elif user_choice.isdigit() and 1 <= int(user_choice) <= 12:
                selected_m = MONTH_INFO[int(user_choice) - 1][2]
    except (EOFError, KeyboardInterrupt):
        selected_m = default_month

    selected_y = default_year
    try:
        user_year = input(f"Enter evaluation year (Sal) [Default: {default_year}]: ").strip()
        if user_year:
            selected_y = user_year
    except (EOFError, KeyboardInterrupt):
        selected_y = default_year
        
    return selected_m, selected_y

def find_reports_folder(month_name):
    # Search priorities:
    # 1. reports/<month_name>
    # 2. reports/
    # 3. گزارشات_ماهانه/<month_name>
    # 4. گزارشات_ماهانه/
    candidates = [
        os.path.join('reports', month_name),
        'reports',
        os.path.join('گزارشات_ماهانه', month_name),
        'گزارشات_ماهانه'
    ]
    for c in candidates:
        if os.path.exists(c):
            files = glob.glob(os.path.join(c, '*.xlsx'))
            files = [f for f in files if not os.path.basename(f).startswith('~$')]
            if files:
                return c, files
                
    # If no files found, default to 'reports'
    os.makedirs('reports', exist_ok=True)
    return 'reports', []

def process_all_reports(master_excel="تهیه کارنامه نواحی.xlsx", selected_month=None, selected_year=None):
    if not os.path.exists(master_excel):
        print(f"❌ Error: Master Excel file '{master_excel}' not found!")
        return False, "شهریور", "1405"
        
    wb_master = openpyxl.load_workbook(master_excel)
    ws_card = wb_master['کارنامه هوشمند']
    
    def_month = str(ws_card['C3'].value or 'شهریور').strip()
    if def_month not in [m[2] for m in MONTH_INFO]:
        def_month = 'شهریور'
        
    def_year = str(ws_card['E3'].value or '1405').strip() if 'E3' in ws_card else '1405'

    if selected_month and selected_year:
        month = selected_month
        year = selected_year
    elif selected_month:
        month = selected_month
        year = def_year
    else:
        month, year = prompt_month_and_year(default_month=def_month, default_year=def_year)
        
    # Update Excel headers
    ws_card['C3'].value = month
    ws_card['E3'].value = str(year)
    
    print("-" * 75)
    print(f"✓ Selected Evaluation Period: Month [{month}] - Year [{year}]")
    
    folder_name, excel_files = find_reports_folder(month)
    print(f"📁 Reports folder: '{folder_name}' | Received Excel files: {len(excel_files)}")
    print("-" * 75)
    
    if not excel_files:
        print(f"⚠️ No employee Excel reports found in '{folder_name}'.")
        print(f"💡 Please copy the monthly report Excel files into folder '{folder_name}'.")
        wb_master.save(master_excel)
        return False, month, str(year)
        
    extracted = {}
    
    for fpath in excel_files:
        fname = os.path.basename(fpath)
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
            detected = match_district_name(fname)
            
            ws_hozori = None
            ws_tavanmand = None
            ws_majazi = None
            ws_khalagh = None
            ws_tolid = None
            
            for sname in wb.sheetnames:
                cn = clean_str(sname)
                if 'حضوری' in cn: ws_hozori = wb[sname]
                elif 'توانمند' in cn: ws_tavanmand = wb[sname]
                elif 'مجازی' in cn or 'لایو' in cn: ws_majazi = wb[sname]
                elif 'خلاق' in cn: ws_khalagh = wb[sname]
                elif 'تولید' in cn: ws_tolid = wb[sname]
                
            if not detected:
                for ws in [ws_hozori, ws_majazi, ws_khalagh, ws_tavanmand]:
                    if ws is None: continue
                    for r in range(2, min(ws.max_row + 1, 25)):
                        for c in [3, 4, 2]:
                            m = match_district_name(ws.cell(r, c).value)
                            if m:
                                detected = m
                                break
                        if detected: break
                    if detected: break
                    
            if not detected:
                print(f"⚠️ Could not detect district for file: '{fname}'")
                continue
                
            m_hoz = extract_sheet_metrics(ws_hozori)
            m_tav = extract_sheet_metrics(ws_tavanmand)
            m_maj = extract_sheet_metrics(ws_majazi)
            m_kha = extract_sheet_metrics(ws_khalagh)
            m_tol = extract_sheet_metrics(ws_tolid)
            
            hoz_tot = m_hoz['people_sum'] + m_tav['people_sum']
            hoz_cls = m_hoz['classes_count'] + m_tav['classes_count']
            metric_hoz = hoz_tot if hoz_tot > 0 else hoz_cls
            
            metric_maj = m_maj['people_sum'] if m_maj['people_sum'] > 0 else m_maj['classes_count']
            metric_kha = m_kha['people_sum'] if m_kha['people_sum'] > 0 else m_kha['classes_count']
            metric_tol = m_tol['people_sum'] if m_tol['people_sum'] > 0 else m_tol['classes_count']
            neshast = 1 if (metric_hoz + metric_maj + metric_kha + metric_tol) > 0 else 0
            
            extracted[detected] = {
                'hozori_total': metric_hoz,
                'majazi': metric_maj,
                'khalagh': metric_kha,
                'tolid': metric_tol,
                'neshast': neshast,
                'details': f"حضوری: {metric_hoz} نفر | مجازی: {metric_maj} نفر | خلاقانه: {metric_kha} نفر | تولیدات: {metric_tol}"
            }
            print(f"✓ [{detected}]: {extracted[detected]['details']}")
            
        except Exception as e:
            print(f"❌ Error processing file '{fname}': {e}")

    print("-" * 75)
    print("💾 Updating metrics into master Excel file...")
    
    ws_rep = wb_master['گزارش عملکرد ماهانه']
    row_map = {}
    for r in range(2, ws_rep.max_row + 1):
        dname = ws_rep.cell(r, 1).value
        if dname:
            row_map[clean_str(dname)] = r
            row_map[clean_no_vav(dname)] = r

    active_count = 0
    inactive_count = 0
    
    for dname in DISTRICTS:
        r = row_map.get(clean_str(dname)) or row_map.get(clean_no_vav(dname))
        if not r:
            continue
            
        if dname in extracted:
            data = extracted[dname]
            ws_rep.cell(row=r, column=2, value=data['hozori_total'])
            ws_rep.cell(row=r, column=3, value=data['majazi'])
            ws_rep.cell(row=r, column=4, value=data['khalagh'])
            ws_rep.cell(row=r, column=5, value=data['tolid'])
            ws_rep.cell(row=r, column=6, value=data['neshast'])
            active_count += 1
        else:
            # District had no file in this month -> zero performance!
            ws_rep.cell(row=r, column=2, value=0)
            ws_rep.cell(row=r, column=3, value=0)
            ws_rep.cell(row=r, column=4, value=0)
            ws_rep.cell(row=r, column=5, value=0)
            ws_rep.cell(row=r, column=6, value=0)
            inactive_count += 1
            print(f"⭕ [{dname}]: گزارشی ارسال نشده (عملکرد ماه {month} صفر ثبت گردید)")

    wb_master.save(master_excel)
    
    # Save a dedicated monthly archive copy
    archive_name = f"کارنامه_نواحی_{month}_{year}.xlsx"
    try:
        shutil.copyfile(master_excel, archive_name)
        print(f"📁 یک نسخه پشتیبان با نام '{archive_name}' ذخیره شد.")
    except Exception:
        pass
        
    print(f"🎉 عملیات ثبت در اکسل کامل شد! {active_count} ناحیه فعال و {inactive_count} ناحیه فاقد فعالیت ثبت گردید.")
    return True, month, str(year)

def generate_all_images_offline(master_excel="تهیه کارنامه نواحی.xlsx", month="شهریور", year="1405"):
    try:
        from image_generator import generate_scorecard_png, generate_dashboard_png
    except Exception as e:
        print("⚠️ ماژول‌های تولید تصویر در دسترس نیستند:", e)
        return
        
    # Main folder and month-dedicated folder
    out_dir_main = "output_cards"
    out_dir_month = os.path.join(out_dir_main, f"{month}_{year}")
    os.makedirs(out_dir_month, exist_ok=True)
    
    print("-" * 75)
    print(f"📸 Generating mobile scorecards for 32 districts ({month} {year}) in '{out_dir_month}'...")
    
    if not os.path.exists(master_excel):
        print(f"❌ فایل '{master_excel}' یافت نشد.")
        return

    wb = openpyxl.load_workbook(master_excel, data_only=True)
    ws_target = wb['پایگاه داده حد انتظار']
    ws_rep = wb['گزارش عملکرد ماهانه']
    
    targets = {}
    for r in range(2, 34):
        dn = ws_target.cell(r, 1).value
        targets[dn] = {
            'branches': ws_target.cell(r, 2).value or 0,
            'hozori': ws_target.cell(r, 3).value or 0,
            'majazi': ws_target.cell(r, 4).value or 0,
            'khalagh': ws_target.cell(r, 5).value or 0,
            'tolid': ws_target.cell(r, 6).value or 0,
            'neshast': ws_target.cell(r, 7).value or 1
        }
        
    actuals = {}
    for r in range(2, ws_rep.max_row + 1):
        dn = ws_rep.cell(r, 1).value
        if dn:
            actuals[dn] = {
                'hozori': ws_rep.cell(r, 2).value or 0,
                'majazi': ws_rep.cell(r, 3).value or 0,
                'khalagh': ws_rep.cell(r, 4).value or 0,
                'tolid': ws_rep.cell(r, 5).value or 0,
                'neshast': ws_rep.cell(r, 6).value or 0
            }
            
    dist_scores = []
    for dn, t in targets.items():
        a = actuals.get(dn, {'hozori': 0, 'majazi': 0, 'khalagh': 0, 'tolid': 0, 'neshast': 0})
        pcts = []
        for k in ['hozori', 'majazi', 'khalagh', 'tolid', 'neshast']:
            tv = t.get(k, 1)
            av = a.get(k, 0)
            pcts.append((av / tv * 100) if tv > 0 else 0)
        sc = sum(pcts) / len(pcts) if pcts else 0
        a['overall_score'] = sc
        dist_scores.append((dn, sc))
        
    dist_scores.sort(key=lambda x: x[1], reverse=True)
    
    rank_map = {}
    current_rank = 1
    for dn, sc in dist_scores:
        if sc > 0:
            rank_map[dn] = str(current_rank)
            current_rank += 1
        else:
            rank_map[dn] = "عدم فعالیت"
    
    count_img = 0
    for dn in DISTRICTS:
        t = targets.get(dn, {})
        a = actuals.get(dn, {'overall_score': 0})
        sc = a.get('overall_score', 0)
        rk = rank_map.get(dn, "عدم فعالیت")
        
        tier = "عالی" if sc >= 100 else ("خوب" if sc >= 75 else ("متوسط" if sc >= 50 else ("ضعیف" if sc > 0 else "فاقد عملکرد")))
        en_name = DISTRICT_EN_NAMES.get(dn, dn)
        
        out_p_month_fa = os.path.join(out_dir_month, f"کارنامه_{dn}.png")
        out_p_month_en = os.path.join(out_dir_month, f"Scorecard_{en_name}.png")
        
        try:
            generate_scorecard_png(dn, t, a, rank=rk, tier=tier, month=month, year=str(year), output_path=out_p_month_fa)
            shutil.copyfile(out_p_month_fa, out_p_month_en)
            count_img += 1
        except Exception as err:
            print(f"Error generating scorecard for {dn}: {err}")

    # Provincial Dashboard
    sum_t_hoz = sum(targets[d]['hozori'] for d in targets)
    sum_t_maj = sum(targets[d]['majazi'] for d in targets)
    sum_t_kha = sum(targets[d]['khalagh'] for d in targets)
    sum_t_tol = sum(targets[d]['tolid'] for d in targets)
    sum_t_nes = sum(targets[d]['neshast'] for d in targets)
    
    sum_a_hoz = sum(actuals.get(d, {}).get('hozori', 0) for d in targets)
    sum_a_maj = sum(actuals.get(d, {}).get('majazi', 0) for d in targets)
    sum_a_kha = sum(actuals.get(d, {}).get('khalagh', 0) for d in targets)
    sum_a_tol = sum(actuals.get(d, {}).get('tolid', 0) for d in targets)
    sum_a_nes = sum(actuals.get(d, {}).get('neshast', 0) for d in targets)
    
    macro_data = [
        ('سواد رسانه حضوری و توانمندسازی (ضریب ۳۱)', int(sum_t_hoz), int(sum_a_hoz)),
        ('سواد رسانه مجازی و لایو (ضریب ۲۱۷)', int(sum_t_maj), int(sum_a_maj)),
        ('اقدامات و ابتکارات خلاقانه (ضریب ۶۲)', int(sum_t_kha), int(sum_a_kha)),
        ('تولیدات رسانه‌ای و محتوایی (ضریب ۳)', int(sum_t_tol), int(sum_a_tol)),
        ('نشست با انجمن مدرسان (۱ نشست)', int(sum_t_nes), int(sum_a_nes))
    ]
    
    active_dists = [x for x in dist_scores if x[1] > 0]
    inactive_dists = [x for x in dist_scores if x[1] == 0]
    
    top5 = [(i+1, active_dists[i][0], active_dists[i][1]) for i in range(min(5, len(active_dists)))]
    all_sorted = active_dists + inactive_dists
    bot5 = [(i+1, all_sorted[len(all_sorted)-1-i][0], all_sorted[len(all_sorted)-1-i][1]) for i in range(min(5, len(all_sorted)))]
    
    avg_sc = sum(x[1] for x in dist_scores) / len(dist_scores) if dist_scores else 0
    top_d = dist_scores[0][0] if dist_scores and dist_scores[0][1] > 0 else "در انتظار"
    rep_c = sum(1 for x in dist_scores if x[1] > 0)
    
    kpi_d = {'avg_score': avg_sc, 'top_district': top_d, 'reported_count': rep_c}
    dash_path_fa = os.path.join(out_dir_month, f"تصویر_داشبورد_مدیریتی_استان_{month}_{year}.png")
    dash_path_en = os.path.join(out_dir_month, f"Dashboard_Provincial_{month}.png")
    
    try:
        generate_dashboard_png(macro_data, top5, bot5, kpi_d, month=month, year=str(year), output_path=dash_path_fa)
        shutil.copyfile(dash_path_fa, dash_path_en)
    except Exception as err:
        print(f"Error generating dashboard: {err}")
    
    print("-" * 75)
    print(f"🎉 32 Mobile Scorecards & Provincial Dashboard saved successfully in:")
    print(f"   📁 '{os.path.abspath(out_dir_month)}'")
    print("=" * 75)
    
    # Automatically pop up the folder in Windows Explorer
    if sys.platform == 'win32':
        try:
            os.system(f'explorer "{os.path.abspath(out_dir_month)}"')
        except Exception:
            pass

if __name__ == '__main__':
    # Parse month and year from command line if passed:
    # e.g.: python run_aggregation.py 4 1405
    # or: python run_aggregation.py tir 1405
    m_arg = None
    y_arg = None
    if len(sys.argv) > 1:
        raw_m = sys.argv[1].strip().lower()
        if raw_m in MONTH_MAP:
            m_arg = MONTH_MAP[raw_m]
        elif raw_m.isdigit() and 1 <= int(raw_m) <= 12:
            m_arg = MONTH_INFO[int(raw_m) - 1][2]
            
    if len(sys.argv) > 2:
        y_arg = sys.argv[2].strip()
        
    success, sel_month, sel_year = process_all_reports(selected_month=m_arg, selected_year=y_arg)
    generate_all_images_offline(month=sel_month, year=sel_year)
