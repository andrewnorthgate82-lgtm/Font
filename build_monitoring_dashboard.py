import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule

def build():
    filename = 'تهیه کارنامه نواحی.xlsx'
    wb = openpyxl.load_workbook(filename)
    
    sheet_title = 'داشبورد مانیتورینگ نواحی'
    if sheet_title in wb.sheetnames:
        del wb[sheet_title]
    
    # Create monitoring sheet as the very first sheet
    ws = wb.create_sheet(title=sheet_title, index=0)
    ws.sheet_view.rightToLeft = True
    ws.views.sheetView[0].showGridLines = True
    
    # Fonts
    FONT_FAMILY = 'IRANSans'
    font_title = Font(name=FONT_FAMILY, size=15, bold=True, color='FFFFFF')
    font_subtitle = Font(name=FONT_FAMILY, size=10, bold=False, color='ECF0F1')
    font_sec_hdr = Font(name=FONT_FAMILY, size=11, bold=True, color='FFFFFF')
    font_grp_hdr = Font(name=FONT_FAMILY, size=10, bold=True, color='FFFFFF')
    font_tbl_hdr = Font(name=FONT_FAMILY, size=9.5, bold=True, color='FFFFFF')
    font_card_lbl = Font(name=FONT_FAMILY, size=9.5, bold=True, color='34495E')
    font_card_sub = Font(name=FONT_FAMILY, size=8.5, bold=False, color='5D6D7E')
    
    font_data = Font(name=FONT_FAMILY, size=9, bold=False, color='181C1F')
    font_data_bold = Font(name=FONT_FAMILY, size=9, bold=True, color='181C1F')
    font_total = Font(name=FONT_FAMILY, size=9.5, bold=True, color='FFFFFF')
    font_max = Font(name=FONT_FAMILY, size=9, bold=True, color='0E6251')
    font_min = Font(name=FONT_FAMILY, size=9, bold=True, color='78281F')
    font_note = Font(name=FONT_FAMILY, size=8.5, italic=True, color='555555')
    
    # Fills
    fill_banner = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid') # Navy
    fill_sub_banner = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid') # Slate
    fill_sec_hdr = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid') # Deep blue
    fill_grp_hdr = PatternFill(start_color='2E4053', end_color='2E4053', fill_type='solid') # Dark slate
    fill_tbl_hdr = PatternFill(start_color='3B5F90', end_color='3B5F90', fill_type='solid') # Steel blue
    fill_zebra_even = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    fill_zebra_odd = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    
    fill_total = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
    fill_avg = PatternFill(start_color='34495E', end_color='34495E', fill_type='solid')
    fill_max = PatternFill(start_color='D4EFDF', end_color='D4EFDF', fill_type='solid')
    fill_min = PatternFill(start_color='FADBD8', end_color='FADBD8', fill_type='solid')
    
    # KPI Card Fills
    fill_card1 = PatternFill(start_color='F2F4F7', end_color='F2F4F7', fill_type='solid')
    fill_card2 = PatternFill(start_color='EBF5FB', end_color='EBF5FB', fill_type='solid')
    fill_card3 = PatternFill(start_color='E8F8F5', end_color='E8F8F5', fill_type='solid')
    fill_card4 = PatternFill(start_color='FEF9E7', end_color='FEF9E7', fill_type='solid')
    fill_card5 = PatternFill(start_color='F5EEF8', end_color='F5EEF8', fill_type='solid')
    
    # Borders
    border_thin_gray = Border(
        left=Side(style='thin', color='D5D8DC'),
        right=Side(style='thin', color='D5D8DC'),
        top=Side(style='thin', color='D5D8DC'),
        bottom=Side(style='thin', color='D5D8DC')
    )
    border_total = Border(
        left=Side(style='thin', color='FFFFFF'),
        right=Side(style='thin', color='FFFFFF'),
        top=Side(style='thin', color='FFFFFF'),
        bottom=Side(style='double', color='FFFFFF')
    )
    
    # Alignments
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    # ==========================================
    # Column Widths (Cols A to W)
    # ==========================================
    widths = {
        'A': 6,   # ردیف
        'B': 21,  # نام ناحیه (شهرستان)
        'C': 11,  # تعداد حوزه
        'D': 11,  # حضوری - حد انتظار
        'E': 11,  # حضوری - عملکرد
        'F': 12,  # حضوری - درصد تحقق
        'G': 12,  # مجازی - حد انتظار
        'H': 11,  # مجازی - عملکرد
        'I': 12,  # مجازی - درصد تحقق
        'J': 11,  # خلاقانه - حد انتظار
        'K': 11,  # خلاقانه - عملکرد
        'L': 12,  # خلاقانه - درصد تحقق
        'M': 11,  # تولیدات - حد انتظار
        'N': 11,  # تولیدات - عملکرد
        'O': 12,  # تولیدات - درصد تحقق
        'P': 10,  # نشست - حد انتظار
        'Q': 10,  # نشست - عملکرد
        'R': 11,  # نشست - درصد تحقق
        'S': 12,  # اعضای انجمن
        'T': 14,  # نمره کل عملکرد
        'U': 11,  # رتبه استانی
        'V': 16,  # سطح ارزیابی
        'W': 14   # وضعیت گزارش
    }
    for col_l, w in widths.items():
        ws.column_dimensions[col_l].width = w
    
    # ==========================================
    # Row 1 & 2: Main Banners
    # ==========================================
    ws.merge_cells('A1:W1')
    cell_a1 = ws['A1']
    cell_a1.value = "سامانه هوشمند مانیتورینگ و ارزیابی جامع عملکرد نواحی و کارمندان نسرا - استان اصفهان"
    cell_a1.font = font_title
    cell_a1.fill = fill_banner
    cell_a1.alignment = align_center
    ws.row_dimensions[1].height = 36
    
    ws.merge_cells('A2:W2')
    cell_a2 = ws['A2']
    cell_a2.value = "سال ارزیابی: ۱۴۰۵ | مبنای سنجش: شاخص‌های ابلاغی به ازای هر حوزه مقاومت | پایش برخط و خودکار متصل به گزارش عملکرد ماهانه"
    cell_a2.font = font_subtitle
    cell_a2.fill = fill_sub_banner
    cell_a2.alignment = align_center
    ws.row_dimensions[2].height = 24
    
    ws.row_dimensions[3].height = 10
    
    # ==========================================
    # Rows 4 to 6: 5 Executive KPI Summary Cards
    # ==========================================
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 20
    
    cards_meta = [
        {
            'cols': ('B', 'D'),
            'fill': fill_card1,
            'top_border': '1B365D',
            'title': "🏢 پوشش تشکیلاتی استان",
            'formula_val': "۳۲ ناحیه | ۲۲۹ حوزه مقاومت",
            'formula_sub': '="مجموع مدرسان هدف: " & TEXT(SUM(\'پایگاه داده حد انتظار\'!$H$2:$H$33),"#,##0") & " نفر"',
            'font_val': Font(name=FONT_FAMILY, size=12.5, bold=True, color='1B365D')
        },
        {
            'cols': ('E', 'H'),
            'fill': fill_card2,
            'top_border': '2980B9',
            'title': "📊 میانگین نمره عملکرد استان",
            'formula_val': "=IFERROR(AVERAGE(T31:T62), 0)",
            'val_num_fmt': '0.0%',
            'formula_sub': '=IF(E5>=1,"تحقق ۱۰۰٪ اهداف استانی",IF(E5>=0.75,"عملکرد بسیار مطلوب",IF(E5>=0.5,"عملکرد متوسط","نیازمند مداخله و پیگیری")))',
            'font_val': Font(name=FONT_FAMILY, size=15, bold=True, color='1A5276')
        },
        {
            'cols': ('I', 'L'),
            'fill': fill_card3,
            'top_border': '27AE60',
            'title': "🏆 پیشتاز عملکرد استان (رتبه ۱)",
            'formula_val': '=IF(MAX(T31:T62)>0, INDEX(B31:B62, MATCH(MAX(T31:T62), T31:T62, 0)), "در انتظار ثبت داده")',
            'formula_sub': '=IF(MAX(T31:T62)>0, "نمره تحقق: " & TEXT(MAX(T31:T62), "0.0%"), "-")',
            'font_val': Font(name=FONT_FAMILY, size=13, bold=True, color='196F3D')
        },
        {
            'cols': ('M', 'P'),
            'fill': fill_card4,
            'top_border': 'D4AC0D',
            'title': "📥 وضعیت دریافت گزارش ماهانه",
            'formula_val': '=COUNTIF(W31:W62, "ثبت شده") & " از ۳۲ ناحیه"',
            'formula_sub': '="نرخ مشارکت: " & TEXT(COUNTIF(W31:W62, "ثبت شده")/32, "0.0%")',
            'font_val': Font(name=FONT_FAMILY, size=13, bold=True, color='7D6608')
        },
        {
            'cols': ('Q', 'W'),
            'fill': fill_card5,
            'top_border': '8E44AD',
            'title': "🎯 تفکیک سطوح کیفی عملکرد نواحی",
            'formula_val': '="عالی: " & COUNTIF(V31:V62, "عالی*") & " | خوب: " & COUNTIF(V31:V62, "خوب*") & " | متوسط: " & COUNTIF(V31:V62, "متوسط*") & " | ضعیف: " & COUNTIF(V31:V62, "ضعیف*")',
            'formula_sub': '="نواحی فاقد گزارش: " & COUNTIF(W31:W62, "فاقد گزارش") & " ناحیه"',
            'font_val': Font(name=FONT_FAMILY, size=11, bold=True, color='5B2C6F')
        }
    ]
    
    for c in cards_meta:
        c1, c2 = c['cols']
        ws.merge_cells(f'{c1}4:{c2}4')
        ws.merge_cells(f'{c1}5:{c2}5')
        ws.merge_cells(f'{c1}6:{c2}6')
        
        top_cell = ws[f'{c1}4']
        top_cell.value = c['title']
        top_cell.font = font_card_lbl
        
        mid_cell = ws[f'{c1}5']
        mid_cell.value = c['formula_val']
        mid_cell.font = c['font_val']
        if 'val_num_fmt' in c:
            mid_cell.number_format = c['val_num_fmt']
            
        bot_cell = ws[f'{c1}6']
        bot_cell.value = c['formula_sub']
        bot_cell.font = font_card_sub
        
        # Style range
        c1_idx = openpyxl.utils.column_index_from_string(c1)
        c2_idx = openpyxl.utils.column_index_from_string(c2)
        for r in range(4, 7):
            for col in range(c1_idx, c2_idx + 1):
                cell = ws.cell(row=r, column=col)
                cell.fill = c['fill']
                cell.alignment = align_center
                
                t_side = Side(style='medium', color=c['top_border']) if r == 4 else Side(style='thin', color='BDC3C7')
                b_side = Side(style='thin', color='BDC3C7')
                l_side = Side(style='thin', color='BDC3C7')
                r_side = Side(style='thin', color='BDC3C7')
                cell.border = Border(top=t_side, bottom=b_side, left=l_side, right=r_side)

    ws.row_dimensions[7].height = 8
    ws.row_dimensions[8].height = 8
    
    # ==========================================
    # Section 1: Macro Provincial Indicator Summary Table (Rows 9 to 18, Cols B to K)
    # ==========================================
    ws.merge_cells('B9:K9')
    s1_title = ws['B9']
    s1_title.value = "۱. ارزیابی و پایش تحقق شاخص‌های کلان در سطح کل استان اصفهان (تجمیعی)"
    s1_title.font = font_sec_hdr
    s1_title.fill = fill_sec_hdr
    s1_title.alignment = align_right
    ws.row_dimensions[9].height = 26
    
    macro_headers = [
        (2, 'عنوان شاخص عملکردی'),
        (3, 'مبنای ابلاغی'),
        (4, 'حد انتظار کل استان'),
        (5, 'عملکرد کل استان'),
        (6, 'انحراف از هدف'),
        (7, 'درصد تحقق استان'),
        (8, 'میانگین تحقق نواحی'),
        (9, 'وضعیت استانی'),
        (10, 'مرجع در گزارش ماهانه'),
        (11, 'توضیحات و اجزا')
    ]
    
    ws.row_dimensions[10].height = 24
    for col_idx, h_text in macro_headers:
        cell = ws.cell(row=10, column=col_idx)
        cell.value = h_text
        cell.font = font_tbl_hdr
        cell.fill = fill_tbl_hdr
        cell.alignment = align_center
        cell.border = border_thin_gray
        
    macro_rows = [
        ("سواد رسانه حضوری و توانمندسازی", "۳۱× به ازای هر حوزه",
         "=SUM('پایگاه داده حد انتظار'!$C$2:$C$33)", "=SUM('گزارش عملکرد ماهانه'!$B$2:$B$33)",
         "=E11-D11", "=IFERROR(E11/D11, 0)", "=AVERAGE(F31:F62)",
         '=IF(G11>=1, "تحقق کامل", IF(G11>=0.75, "مطلوب", IF(G11>=0.5, "متوسط", "نیازمند جهش")))',
         "شیت حضوری + توانمندسازی", "مجموع کارگاه‌ها و نشست‌های حضوری"),
        
        ("سواد رسانه مجازی و لایو", "۲۱۷× به ازای هر حوزه",
         "=SUM('پایگاه داده حد انتظار'!$D$2:$D$33)", "=SUM('گزارش عملکرد ماهانه'!$C$2:$C$33)",
         "=E12-D12", "=IFERROR(E12/D12, 0)", "=AVERAGE(I31:I62)",
         '=IF(G12>=1, "تحقق کامل", IF(G12>=0.75, "مطلوب", IF(G12>=0.5, "متوسط", "نیازمند جهش")))',
         "شیت سوادرسانه مجازی", "پخش زنده، وبینار و آموزش مجازی"),
        
        ("اقدامات و ابتکارات خلاقانه", "۶۲× به ازای هر حوزه",
         "=SUM('پایگاه داده حد انتظار'!$E$2:$E$33)", "=SUM('گزارش عملکرد ماهانه'!$D$2:$D$33)",
         "=E13-D13", "=IFERROR(E13/D13, 0)", "=AVERAGE(L31:L62)",
         '=IF(G13>=1, "تحقق کامل", IF(G13>=0.75, "مطلوب", IF(G13>=0.5, "متوسط", "نیازمند جهش")))',
         "شیت سواد رسانه خلاقانه", "ایستگاه، نمایشگاه، رادیو، تبلیغات"),
        
        ("تولیدات رسانه‌ای و محتوایی", "۳× به ازای هر حوزه",
         "=SUM('پایگاه داده حد انتظار'!$F$2:$F$33)", "=SUM('گزارش عملکرد ماهانه'!$E$2:$E$33)",
         "=E14-D14", "=IFERROR(E14/D14, 0)", "=AVERAGE(O31:O62)",
         '=IF(G14>=1, "تحقق کامل", IF(G14>=0.75, "مطلوب", IF(G14>=0.5, "متوسط", "نیازمند جهش")))',
         "شیت تولیدات", "کلیپ، اینفوگرافیک، موشن، پوستر"),
        
        ("نشست دبیر با انجمن مدرسان", "۱ نشست به ازای ناحیه",
         "=SUM('پایگاه داده حد انتظار'!$G$2:$G$33)", "=SUM('گزارش عملکرد ماهانه'!$F$2:$F$33)",
         "=E15-D15", "=IFERROR(E15/D15, 0)", "=AVERAGE(R31:R62)",
         '=IF(G15>=1, "تحقق کامل", IF(G15>=0.75, "مطلوب", IF(G15>=0.5, "متوسط", "نیازمند جهش")))',
         "جلسات هماهنگی ناحیه", "نشست رسمی دبیر با مدرسان"),
        
        ("اعضای انجمن مدرسان", "۲۲ نفر به ازای ناحیه",
         "=SUM('پایگاه داده حد انتظار'!$H$2:$H$33)", "-",
         "-", "-", "-",
         "ظرفیت‌سازی", "پایگاه حد انتظار", "مجموع ۷۰۴ عضو فعال استانی")
    ]
    
    for r_idx, row_data in enumerate(macro_rows, start=11):
        ws.row_dimensions[r_idx].height = 21
        is_odd = (r_idx % 2 == 1)
        row_fill = fill_zebra_odd if is_odd else fill_zebra_even
        
        for c_idx, val in enumerate(row_data, start=2):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.value = val
            cell.fill = row_fill
            cell.border = border_thin_gray
            cell.font = font_data
            
            if c_idx in [2, 10, 11]:
                cell.alignment = align_right
            else:
                cell.alignment = align_center
                
            # Number formats
            if c_idx in [4, 5, 6]:
                if isinstance(val, str) and val.startswith('='):
                    cell.number_format = '#,##0'
            elif c_idx in [7, 8]:
                if isinstance(val, str) and val.startswith('='):
                    cell.number_format = '0.0%'
                    cell.font = font_data_bold
                    
    # Row 17: Macro Totals
    ws.row_dimensions[17].height = 24
    total_data = [
        "مجموع فعالیت‌های کلان استان", "-",
        "=SUM(D11:D15)", "=SUM(E11:E15)", "=E17-D17",
        "=IFERROR(E17/D17, 0)", "=AVERAGE(G11:G15)",
        '=IF(G17>=1, "تحقق کامل", IF(G17>=0.75, "مطلوب", IF(G17>=0.5, "متوسط", "نیازمند جهش")))',
        "۵ شاخص اصلی", "مجموع ۷۱,۷۰۹ اقدام استانی"
    ]
    for c_idx, val in enumerate(total_data, start=2):
        cell = ws.cell(row=17, column=c_idx)
        cell.value = val
        cell.fill = fill_total
        cell.font = font_total
        cell.border = border_total
        if c_idx in [2, 10, 11]:
            cell.alignment = align_right
        else:
            cell.alignment = align_center
        if c_idx in [4, 5, 6]:
            cell.number_format = '#,##0'
        elif c_idx in [7, 8]:
            cell.number_format = '0.0%'

    # Embed Bar Chart for Macro Fulfillment (Cols M to W, Rows 9 to 18)
    chart = BarChart()
    chart.type = 'col'
    chart.style = 10
    chart.title = 'درصد تحقق شاخص‌های کلان استان اصفهان'
    chart.y_axis.title = 'درصد تحقق'
    chart.y_axis.number_format = '0%'
    chart.legend = None
    chart.width = 16.5
    chart.height = 7.5
    
    data_ref = Reference(ws, min_col=7, min_row=10, max_row=15)
    cats_ref = Reference(ws, min_col=2, min_row=11, max_row=15)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws.add_chart(chart, 'M9')
    
    ws.row_dimensions[18].height = 10
    ws.row_dimensions[19].height = 10
    
    # ==========================================
    # Section 2: Top 5 & Bottom 5 Leaderboards (Rows 20 to 26)
    # Both tables 11 columns wide: Top 5 = B to L, Bottom 5 = M to W
    # ==========================================
    # Top 5 Performers (Cols B to L)
    ws.merge_cells('B20:L20')
    top5_title = ws['B20']
    top5_title.value = "🏆 دیده‌بان برترین‌های عملکرد (۵ ناحیه پیشتاز استان اصفهان)"
    top5_title.font = font_sec_hdr
    top5_title.fill = PatternFill(start_color='117864', end_color='117864', fill_type='solid') # Emerald
    top5_title.alignment = align_right
    ws.row_dimensions[20].height = 24
    
    top5_headers = [
        (2, 'رتبه'), (3, 'نام ناحیه (شهرستان)'), (4, 'تعداد حوزه'),
        (5, 'نمره کل عملکرد'), (6, 'سطح ارزیابی'), (7, 'وضعیت ثبت'),
        (8, 'حضوری %'), (9, 'مجازی %'), (10, 'خلاقانه %'), (11, 'تولیدات %'), (12, 'نشست %')
    ]
    ws.row_dimensions[21].height = 22
    for c_idx, h_text in top5_headers:
        cell = ws.cell(row=21, column=c_idx)
        cell.value = h_text
        cell.font = font_tbl_hdr
        cell.fill = PatternFill(start_color='16A085', end_color='16A085', fill_type='solid')
        cell.alignment = align_center
        cell.border = border_thin_gray
        
    for k in range(1, 6):
        r_idx = 21 + k
        ws.row_dimensions[r_idx].height = 20
        row_fill = fill_zebra_odd if (k % 2 == 1) else fill_zebra_even
        
        row_cells = [
            (2, k, '@'),
            (3, f'=IF(MAX($T$31:$T$62)>0, INDEX($B$31:$B$62, MATCH(LARGE($T$31:$T$62, {k}), $T$31:$T$62, 0)), "در انتظار ثبت عملکرد")', '@'),
            (4, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(C{r_idx}, $B$31:$C$62, 2, FALSE), "-")', '#,##0'),
            (5, f'=IF(MAX($T$31:$T$62)>0, LARGE($T$31:$T$62, {k}), 0)', '0.0%'),
            (6, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(C{r_idx}, $B$31:$V$62, 21, FALSE), "-")', '@'),
            (7, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(C{r_idx}, $B$31:$W$62, 22, FALSE), "-")', '@'),
            (8, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(C{r_idx}, $B$31:$F$62, 5, FALSE), "-")', '0.0%'),
            (9, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(C{r_idx}, $B$31:$I$62, 8, FALSE), "-")', '0.0%'),
            (10, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(C{r_idx}, $B$31:$L$62, 11, FALSE), "-")', '0.0%'),
            (11, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(C{r_idx}, $B$31:$O$62, 14, FALSE), "-")', '0.0%'),
            (12, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(C{r_idx}, $B$31:$R$62, 17, FALSE), "-")', '0.0%')
        ]
        for c_idx, f_val, num_fmt in row_cells:
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.value = f_val
            cell.fill = row_fill
            cell.border = border_thin_gray
            cell.font = font_data_bold if c_idx in [2, 5] else font_data
            cell.alignment = align_right if c_idx == 3 else align_center
            cell.number_format = num_fmt

    # Bottom 5 / Watchlist (Cols M to W)
    ws.merge_cells('M20:W20')
    bot5_title = ws['M20']
    bot5_title.value = "⚠️ دیده‌بان پیگیری و تقویت (۵ ناحیه با کمترین تحقق / نیازمند توجه)"
    bot5_title.font = font_sec_hdr
    bot5_title.fill = PatternFill(start_color='C0392B', end_color='C0392B', fill_type='solid') # Red
    bot5_title.alignment = align_right
    
    bot5_headers = [
        (13, 'اولویت'), (14, 'نام ناحیه (شهرستان)'), (15, 'تعداد حوزه'),
        (16, 'نمره کل عملکرد'), (17, 'سطح ارزیابی'), (18, 'وضعیت ثبت'),
        (19, 'حضوری %'), (20, 'مجازی %'), (21, 'خلاقانه %'), (22, 'اقدام پیشنهادی'), (23, 'پیوند کارنامه')
    ]
    for c_idx, h_text in bot5_headers:
        cell = ws.cell(row=21, column=c_idx)
        cell.value = h_text
        cell.font = font_tbl_hdr
        cell.fill = PatternFill(start_color='E74C3C', end_color='E74C3C', fill_type='solid')
        cell.alignment = align_center
        cell.border = border_thin_gray

    for k in range(1, 6):
        r_idx = 21 + k
        row_fill = fill_zebra_odd if (k % 2 == 1) else fill_zebra_even
        
        row_cells = [
            (13, k, '@'),
            (14, f'=IF(MAX($T$31:$T$62)>0, INDEX($B$31:$B$62, MATCH(SMALL($T$31:$T$62, {k}), $T$31:$T$62, 0)), "در انتظار ثبت عملکرد")', '@'),
            (15, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(N{r_idx}, $B$31:$C$62, 2, FALSE), "-")', '#,##0'),
            (16, f'=IF(MAX($T$31:$T$62)>0, SMALL($T$31:$T$62, {k}), 0)', '0.0%'),
            (17, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(N{r_idx}, $B$31:$V$62, 21, FALSE), "-")', '@'),
            (18, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(N{r_idx}, $B$31:$W$62, 22, FALSE), "-")', '@'),
            (19, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(N{r_idx}, $B$31:$F$62, 5, FALSE), "-")', '0.0%'),
            (20, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(N{r_idx}, $B$31:$I$62, 8, FALSE), "-")', '0.0%'),
            (21, f'=IF(MAX($T$31:$T$62)>0, VLOOKUP(N{r_idx}, $B$31:$L$62, 11, FALSE), "-")', '0.0%'),
            (22, f'=IF(MAX($T$31:$T$62)>0, "پیگیری کسری عملکرد و کارگاه‌ها", "-")', '@'),
            (23, '=HYPERLINK("#\'کارنامه هوشمند\'!A1", "مشاهده کارنامه ↗")', '@')
        ]
        for c_idx, f_val, num_fmt in row_cells:
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.value = f_val
            cell.fill = row_fill
            cell.border = border_thin_gray
            cell.font = font_data_bold if c_idx in [13, 16] else font_data
            cell.alignment = align_right if c_idx in [14, 22] else align_center
            cell.number_format = num_fmt

    ws.row_dimensions[27].height = 12
    
    # ==========================================
    # Section 3: Master 32-District Performance Table (Rows 28 to 66)
    # ==========================================
    ws.merge_cells('A28:W28')
    s3_title = ws['A28']
    s3_title.value = "۲. جدول جامع رصد و مانیتورینگ عملکرد ۳۲ ناحیه و کارمند استان اصفهان (تفکیک شاخص‌ها و حوزه‌ها)"
    s3_title.font = font_sec_hdr
    s3_title.fill = fill_sec_hdr
    s3_title.alignment = align_right
    ws.row_dimensions[28].height = 26
    
    # Row 29: Group Headers
    ws.row_dimensions[29].height = 24
    groups = [
        ('A29:C29', "مشخصات کلی ناحیه", fill_grp_hdr),
        ('D29:F29', "۱. حضوری و توانمندسازی (۳۱×)", PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')),
        ('G29:I29', "۲. مجازی و لایو (۲۱۷×)", PatternFill(start_color='2874A6', end_color='2874A6', fill_type='solid')),
        ('J29:L29', "۳. اقدامات خلاقانه (۶۲×)", PatternFill(start_color='117864', end_color='117864', fill_type='solid')),
        ('M29:O29', "۴. تولیدات رسانه‌ای (۳×)", PatternFill(start_color='7D6608', end_color='7D6608', fill_type='solid')),
        ('P29:R29', "۵. نشست انجمن (۱×)", PatternFill(start_color='512E5F', end_color='512E5F', fill_type='solid')),
        ('S29:S29', "۶. اعضا", PatternFill(start_color='34495E', end_color='34495E', fill_type='solid')),
        ('T29:W29', "ارزیابی نهایی و رتبه‌بندی استانی", PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid'))
    ]
    for rng, g_text, g_fill in groups:
        if ':' in rng:
            ws.merge_cells(rng)
            c_start = rng.split(':')[0]
        else:
            c_start = rng
        top_cell = ws[c_start]
        top_cell.value = g_text
        top_cell.font = font_grp_hdr
        top_cell.alignment = align_center
        
        # style all cells in group header
        if ':' in rng:
            parts = rng.split(':')
            c1_i = openpyxl.utils.column_index_from_string(parts[0][:len(parts[0])-2])
            c2_i = openpyxl.utils.column_index_from_string(parts[1][:len(parts[1])-2])
        else:
            c1_i = c2_i = openpyxl.utils.column_index_from_string(rng[:len(rng)-2])
            
        for c_i in range(c1_i, c2_i + 1):
            cell = ws.cell(row=29, column=c_i)
            cell.fill = g_fill
            cell.border = border_thin_gray
            
    # Row 30: Detailed Column Headers
    ws.row_dimensions[30].height = 24
    col_headers = [
        (1, 'ردیف'),
        (2, 'نام ناحیه (شهرستان)'),
        (3, 'تعداد حوزه'),
        (4, 'حد انتظار'),
        (5, 'عملکرد'),
        (6, 'درصد تحقق'),
        (7, 'حد انتظار'),
        (8, 'عملکرد'),
        (9, 'درصد تحقق'),
        (10, 'حد انتظار'),
        (11, 'عملکرد'),
        (12, 'درصد تحقق'),
        (13, 'حد انتظار'),
        (14, 'عملکرد'),
        (15, 'درصد تحقق'),
        (16, 'حد انتظار'),
        (17, 'عملکرد'),
        (18, 'درصد تحقق'),
        (19, 'اعضای انجمن'),
        (20, 'نمره کل عملکرد'),
        (21, 'رتبه استانی'),
        (22, 'سطح ارزیابی'),
        (23, 'وضعیت گزارش')
    ]
    for col_idx, h_text in col_headers:
        cell = ws.cell(row=30, column=col_idx)
        cell.value = h_text
        cell.font = font_tbl_hdr
        cell.fill = fill_tbl_hdr
        cell.alignment = align_center
        cell.border = border_thin_gray

    # District Names (All 32 districts)
    districts = [
        'آران و بیدگل', 'امام حسین(ع)', 'امام رضا(ع)', 'امام صادق(ع)', 'امام علی(ع)',
        'اردستان', 'برخوار', 'بویین و میاندشت', 'تیران و کرون', 'جرقویه',
        'چادگان', 'خمینی شهر', 'خوانسار', 'خور و بیابانک', 'درچه',
        'دهاقان', 'سمیرم', 'شاهین شهر', 'شهرضا', 'فریدن',
        'فریدون شهر', 'فلاورجان', 'کاشان', 'کوهپایه', 'گلپایگان',
        'لنجان', 'مبارکه', 'نایین', 'نجف آباد', 'نطنز',
        'ورزنه', 'هرند'
    ]
    
    for i, d_name in enumerate(districts, start=1):
        r = 30 + i
        ws.row_dimensions[r].height = 20.5
        is_odd = (i % 2 == 1)
        row_fill = fill_zebra_odd if is_odd else fill_zebra_even
        
        row_formulas = [
            (1, i, '@', align_center, font_data),
            (2, d_name, '@', align_right, font_data_bold),
            (3, f"=VLOOKUP($B{r}, 'پایگاه داده حد انتظار'!$A$2:$H$33, 2, FALSE)", '#,##0', align_center, font_data_bold),
            
            # 1. حضوری و توانمندسازی (۳۱×)
            (4, f"=VLOOKUP($B{r}, 'پایگاه داده حد انتظار'!$A$2:$H$33, 3, FALSE)", '#,##0', align_center, font_data),
            (5, f"=IFERROR(VLOOKUP($B{r}, 'گزارش عملکرد ماهانه'!$A$2:$F$33, 2, FALSE), 0)", '#,##0', align_center, font_data),
            (6, f"=IFERROR(E{r}/D{r}, 0)", '0.0%', align_center, font_data_bold),
            
            # 2. مجازی و لایو (۲۱۷×)
            (7, f"=VLOOKUP($B{r}, 'پایگاه داده حد انتظار'!$A$2:$H$33, 4, FALSE)", '#,##0', align_center, font_data),
            (8, f"=IFERROR(VLOOKUP($B{r}, 'گزارش عملکرد ماهانه'!$A$2:$F$33, 3, FALSE), 0)", '#,##0', align_center, font_data),
            (9, f"=IFERROR(H{r}/G{r}, 0)", '0.0%', align_center, font_data_bold),
            
            # 3. اقدامات خلاقانه (۶۲×)
            (10, f"=VLOOKUP($B{r}, 'پایگاه داده حد انتظار'!$A$2:$H$33, 5, FALSE)", '#,##0', align_center, font_data),
            (11, f"=IFERROR(VLOOKUP($B{r}, 'گزارش عملکرد ماهانه'!$A$2:$F$33, 4, FALSE), 0)", '#,##0', align_center, font_data),
            (12, f"=IFERROR(K{r}/J{r}, 0)", '0.0%', align_center, font_data_bold),
            
            # 4. تولیدات رسانه‌ای (۳×)
            (13, f"=VLOOKUP($B{r}, 'پایگاه داده حد انتظار'!$A$2:$H$33, 6, FALSE)", '#,##0', align_center, font_data),
            (14, f"=IFERROR(VLOOKUP($B{r}, 'گزارش عملکرد ماهانه'!$A$2:$F$33, 5, FALSE), 0)", '#,##0', align_center, font_data),
            (15, f"=IFERROR(N{r}/M{r}, 0)", '0.0%', align_center, font_data_bold),
            
            # 5. نشست انجمن (۱×)
            (16, f"=VLOOKUP($B{r}, 'پایگاه داده حد انتظار'!$A$2:$H$33, 7, FALSE)", '#,##0', align_center, font_data),
            (17, f"=IFERROR(VLOOKUP($B{r}, 'گزارش عملکرد ماهانه'!$A$2:$F$33, 6, FALSE), 0)", '#,##0', align_center, font_data),
            (18, f"=IFERROR(Q{r}/P{r}, 0)", '0.0%', align_center, font_data_bold),
            
            # 6. اعضای انجمن (۲۲ نفر)
            (19, f"=VLOOKUP($B{r}, 'پایگاه داده حد انتظار'!$A$2:$H$33, 8, FALSE)", '#,##0', align_center, font_data),
            
            # Final Metrics
            (20, f"=AVERAGE(F{r}, I{r}, L{r}, O{r}, R{r})", '0.0%', align_center, Font(name=FONT_FAMILY, size=9.5, bold=True, color='1B365D')),
            (21, f'=IF(T{r}>0, RANK(T{r}, $T$31:$T$62), "-")', '@', align_center, font_data_bold),
            (22, f'=IF(W{r}="فاقد گزارش", "ثبت نشده", IF(T{r}>=1, "عالی (۱۰۰٪+)", IF(T{r}>=0.75, "خوب (۷۵-۹۹٪)", IF(T{r}>=0.5, "متوسط (۵۰-۷۴٪)", "ضعیف (زیر ۵۰٪)"))))', '@', align_center, font_data),
            (23, f'=IF((E{r}+H{r}+K{r}+N{r}+Q{r})>0, "ثبت شده", "فاقد گزارش")', '@', align_center, font_data)
        ]
        
        for c_idx, f_val, num_fmt, aln, fnt in row_formulas:
            cell = ws.cell(row=r, column=c_idx)
            cell.value = f_val
            cell.fill = row_fill
            cell.border = border_thin_gray
            cell.font = fnt
            cell.alignment = aln
            cell.number_format = num_fmt

    # Summary Rows (Rows 63 to 66)
    # Row 63: مجموع کل استان
    ws.row_dimensions[63].height = 24
    row63_data = [
        (1, "-", '@', align_center),
        (2, "مجموع کل استان اصفهان", '@', align_right),
        (3, "=SUM(C31:C62)", '#,##0', align_center),
        (4, "=SUM(D31:D62)", '#,##0', align_center),
        (5, "=SUM(E31:E62)", '#,##0', align_center),
        (6, "=IFERROR(E63/D63, 0)", '0.0%', align_center),
        (7, "=SUM(G31:G62)", '#,##0', align_center),
        (8, "=SUM(H31:H62)", '#,##0', align_center),
        (9, "=IFERROR(H63/G63, 0)", '0.0%', align_center),
        (10, "=SUM(J31:J62)", '#,##0', align_center),
        (11, "=SUM(K31:K62)", '#,##0', align_center),
        (12, "=IFERROR(K63/J63, 0)", '0.0%', align_center),
        (13, "=SUM(M31:M62)", '#,##0', align_center),
        (14, "=SUM(N31:N62)", '#,##0', align_center),
        (15, "=IFERROR(N63/M63, 0)", '0.0%', align_center),
        (16, "=SUM(P31:P62)", '#,##0', align_center),
        (17, "=SUM(Q31:Q62)", '#,##0', align_center),
        (18, "=IFERROR(Q63/P63, 0)", '0.0%', align_center),
        (19, "=SUM(S31:S62)", '#,##0', align_center),
        (20, "=AVERAGE(T31:T62)", '0.0%', align_center),
        (21, "-", '@', align_center),
        (22, '=IF(T63>=1, "عالی", IF(T63>=0.75, "خوب", IF(T63>=0.5, "متوسط", "نیازمند جهش")))', '@', align_center),
        (23, '=COUNTIF(W31:W62, "ثبت شده") & " از ۳۲ ناحیه"', '@', align_center)
    ]
    for c_idx, f_val, num_fmt, aln in row63_data:
        cell = ws.cell(row=63, column=c_idx)
        cell.value = f_val
        cell.fill = fill_total
        cell.font = font_total
        cell.border = border_total
        cell.alignment = aln
        cell.number_format = num_fmt

    # Row 64: میانگین عملکرد نواحی
    ws.row_dimensions[64].height = 22
    row64_data = [
        (1, "-", '@', align_center),
        (2, "میانگین عملکرد هر ناحیه", '@', align_right),
        (3, "=AVERAGE(C31:C62)", '#,##0.0', align_center),
        (4, "=AVERAGE(D31:D62)", '#,##0.0', align_center),
        (5, "=AVERAGE(E31:E62)", '#,##0.0', align_center),
        (6, "=AVERAGE(F31:F62)", '0.0%', align_center),
        (7, "=AVERAGE(G31:G62)", '#,##0.0', align_center),
        (8, "=AVERAGE(H31:H62)", '#,##0.0', align_center),
        (9, "=AVERAGE(I31:I62)", '0.0%', align_center),
        (10, "=AVERAGE(J31:J62)", '#,##0.0', align_center),
        (11, "=AVERAGE(K31:K62)", '#,##0.0', align_center),
        (12, "=AVERAGE(L31:L62)", '0.0%', align_center),
        (13, "=AVERAGE(M31:M62)", '#,##0.0', align_center),
        (14, "=AVERAGE(N31:N62)", '#,##0.0', align_center),
        (15, "=AVERAGE(O31:O62)", '0.0%', align_center),
        (16, "=AVERAGE(P31:P62)", '#,##0.0', align_center),
        (17, "=AVERAGE(Q31:Q62)", '#,##0.0', align_center),
        (18, "=AVERAGE(R31:R62)", '0.0%', align_center),
        (19, "=AVERAGE(S31:S62)", '#,##0.0', align_center),
        (20, "=AVERAGE(T31:T62)", '0.0%', align_center),
        (21, "-", '@', align_center),
        (22, "-", '@', align_center),
        (23, "-", '@', align_center)
    ]
    for c_idx, f_val, num_fmt, aln in row64_data:
        cell = ws.cell(row=64, column=c_idx)
        cell.value = f_val
        cell.fill = fill_avg
        cell.font = font_total
        cell.border = border_thin_gray
        cell.alignment = aln
        cell.number_format = num_fmt

    # Row 65: حداکثر عملکرد (MAX)
    ws.row_dimensions[65].height = 21
    row65_data = [
        (1, "-", '@', align_center),
        (2, "بالاترین رکورد در استان (حداکثر)", '@', align_right),
        (3, "=MAX(C31:C62)", '#,##0', align_center),
        (4, "=MAX(D31:D62)", '#,##0', align_center),
        (5, "=MAX(E31:E62)", '#,##0', align_center),
        (6, "=MAX(F31:F62)", '0.0%', align_center),
        (7, "=MAX(G31:G62)", '#,##0', align_center),
        (8, "=MAX(H31:H62)", '#,##0', align_center),
        (9, "=MAX(I31:I62)", '0.0%', align_center),
        (10, "=MAX(J31:J62)", '#,##0', align_center),
        (11, "=MAX(K31:K62)", '#,##0', align_center),
        (12, "=MAX(L31:L62)", '0.0%', align_center),
        (13, "=MAX(M31:M62)", '#,##0', align_center),
        (14, "=MAX(N31:N62)", '#,##0', align_center),
        (15, "=MAX(O31:O62)", '0.0%', align_center),
        (16, "=MAX(P31:P62)", '#,##0', align_center),
        (17, "=MAX(Q31:Q62)", '#,##0', align_center),
        (18, "=MAX(R31:R62)", '0.0%', align_center),
        (19, "=MAX(S31:S62)", '#,##0', align_center),
        (20, "=MAX(T31:T62)", '0.0%', align_center),
        (21, "-", '@', align_center),
        (22, "-", '@', align_center),
        (23, "-", '@', align_center)
    ]
    for c_idx, f_val, num_fmt, aln in row65_data:
        cell = ws.cell(row=65, column=c_idx)
        cell.value = f_val
        cell.fill = fill_max
        cell.font = font_max
        cell.border = border_thin_gray
        cell.alignment = aln
        cell.number_format = num_fmt

    # Row 66: حداقل عملکرد (MIN)
    ws.row_dimensions[66].height = 21
    row66_data = [
        (1, "-", '@', align_center),
        (2, "کمترین رکورد در استان (حداقل)", '@', align_right),
        (3, "=MIN(C31:C62)", '#,##0', align_center),
        (4, "=MIN(D31:D62)", '#,##0', align_center),
        (5, "=MIN(E31:E62)", '#,##0', align_center),
        (6, "=MIN(F31:F62)", '0.0%', align_center),
        (7, "=MIN(G31:G62)", '#,##0', align_center),
        (8, "=MIN(H31:H62)", '#,##0', align_center),
        (9, "=MIN(I31:I62)", '0.0%', align_center),
        (10, "=MIN(J31:J62)", '#,##0', align_center),
        (11, "=MIN(K31:K62)", '#,##0', align_center),
        (12, "=MIN(L31:L62)", '0.0%', align_center),
        (13, "=MIN(M31:M62)", '#,##0', align_center),
        (14, "=MIN(N31:N62)", '#,##0', align_center),
        (15, "=MIN(O31:O62)", '0.0%', align_center),
        (16, "=MIN(P31:P62)", '#,##0', align_center),
        (17, "=MIN(Q31:Q62)", '#,##0', align_center),
        (18, "=MIN(R31:R62)", '0.0%', align_center),
        (19, "=MIN(S31:S62)", '#,##0', align_center),
        (20, "=MIN(T31:T62)", '0.0%', align_center),
        (21, "-", '@', align_center),
        (22, "-", '@', align_center),
        (23, "-", '@', align_center)
    ]
    for c_idx, f_val, num_fmt, aln in row66_data:
        cell = ws.cell(row=66, column=c_idx)
        cell.value = f_val
        cell.fill = fill_min
        cell.font = font_min
        cell.border = border_thin_gray
        cell.alignment = aln
        cell.number_format = num_fmt

    ws.row_dimensions[67].height = 12

    # ==========================================
    # Section 4: Workflow Instructions & Direct Links (Rows 68 to 73)
    # ==========================================
    ws.merge_cells('A68:W68')
    ws['A68'].value = "📌 راهنمای گردش اطلاعات و پیوندهای سریع میان شیت‌ها:"
    ws['A68'].font = Font(name=FONT_FAMILY, size=10.5, bold=True, color='1B365D')
    ws['A68'].alignment = align_right
    ws.row_dimensions[68].height = 22
    
    notes = [
        "۱. نحوه ورود داده‌ها: گزارش‌های ارسالی ماهانه نواحی را تجمیع و در شیت «گزارش عملکرد ماهانه» درج فرمایید؛ تمام محاسبات، درصدها، رتبه‌ها و نمودار این داشبورد به صورت آنی و خودکار بروزرسانی خواهند شد.",
        "۲. مبنای محاسبه شاخص‌ها: ملاک محاسبه عملکرد در شیت‌های حضوری، توانمندسازی، مجازی و خلاقانه، «مجموع تعداد نفرات / مخاطبان» ثبت‌شده در ستون تعداد نفرات کلاس‌ها است (در شاخص ۱: مجموع نفرات شیت‌های حضوری و توانمندسازی گردان).",
        "۳. صدور کارنامه تک‌برگی: جهت مشاهده یا چاپ کارنامه رسمی هر شهرستان به صورت انفرادی، به شیت «کارنامه هوشمند» مراجعه و نام شهرستان را از منوی کشویی انتخاب فرمایید."
    ]
    for idx, note_txt in enumerate(notes, start=69):
        ws.merge_cells(f'A{idx}:W{idx}')
        cell = ws[f'A{idx}']
        cell.value = note_txt
        cell.font = font_note
        cell.alignment = align_right
        ws.row_dimensions[idx].height = 20

    # Links row
    ws.row_dimensions[73].height = 26
    ws.merge_cells('B73:E73')
    lnk1 = ws['B73']
    lnk1.value = '=HYPERLINK("#\'کارنامه هوشمند\'!A1", "📋 ورود به شیت صدور کارنامه هوشمند نواحی ↗")'
    lnk1.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color='1B4F72', underline='single')
    lnk1.fill = fill_card2
    lnk1.alignment = align_center
    lnk1.border = border_thin_gray
    
    ws.merge_cells('F73:I73')
    lnk2 = ws['F73']
    lnk2.value = '=HYPERLINK("#\'گزارش عملکرد ماهانه\'!A1", "✏️ ورود به شیت ثبت عملکرد ماهانه نواحی ↗")'
    lnk2.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color='196F3D', underline='single')
    lnk2.fill = fill_card3
    lnk2.alignment = align_center
    lnk2.border = border_thin_gray
    
    ws.merge_cells('J73:M73')
    lnk3 = ws['J73']
    lnk3.value = '=HYPERLINK("#\'پایگاه داده حد انتظار\'!A1", "⚙️ ورود به پایگاه داده و حدانتظار حوزه‌ها ↗")'
    lnk3.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color='7D6608', underline='single')
    lnk3.fill = fill_card4
    lnk3.alignment = align_center
    lnk3.border = border_thin_gray

    # ==========================================
    # Conditional Formatting for Achievement Percentages
    # ==========================================
    cf_ranges = ['F31:F62', 'I31:I62', 'L31:L62', 'O31:O62', 'R31:R62', 'T31:T62']
    
    rule_green = CellIsRule(operator='greaterThanOrEqual', formula=['1'], stopIfTrue=False,
                            fill=PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid'),
                            font=Font(name=FONT_FAMILY, color='155724', bold=True))
    rule_blue = CellIsRule(operator='between', formula=['0.75', '0.9999'], stopIfTrue=False,
                           fill=PatternFill(start_color='D1ECF1', end_color='D1ECF1', fill_type='solid'),
                           font=Font(name=FONT_FAMILY, color='0C5460', bold=True))
    rule_yellow = CellIsRule(operator='between', formula=['0.5', '0.7499'], stopIfTrue=False,
                            fill=PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid'),
                            font=Font(name=FONT_FAMILY, color='856404', bold=True))
    rule_red = CellIsRule(operator='between', formula=['0.0001', '0.4999'], stopIfTrue=False,
                         fill=PatternFill(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid'),
                         font=Font(name=FONT_FAMILY, color='721C24', bold=True))

    for rng in cf_ranges:
        ws.conditional_formatting.add(rng, rule_green)
        ws.conditional_formatting.add(rng, rule_blue)
        ws.conditional_formatting.add(rng, rule_yellow)
        ws.conditional_formatting.add(rng, rule_red)

    # ==========================================
    # Enhance 'کارنامه هوشمند' sheet with return link & extra info
    # ==========================================
    ws_k = wb['کارنامه هوشمند']
    ws_k.sheet_view.rightToLeft = True
    
    # Navigation link at B1
    ws_k['B1'] = '=HYPERLINK("#\'داشبورد مانیتورینگ نواحی\'!A1", "📊 بازگشت به داشبورد مانیتورینگ جامع نواحی ↗")'
    ws_k['B1'].font = Font(name=FONT_FAMILY, size=9.5, bold=True, color='1B4F72', underline='single')
    ws_k['B1'].alignment = align_right
    
    # Format B11:C11
    ws_k['B11'].font = Font(name=FONT_FAMILY, size=10, bold=True, color='1B365D')
    ws_k['B11'].border = border_thin_gray
    ws_k['B11'].fill = PatternFill(start_color='EAECEE', end_color='EAECEE', fill_type='solid')
    ws_k['C11'].font = Font(name=FONT_FAMILY, size=11, bold=True, color='1B365D')
    ws_k['C11'].border = border_thin_gray
    ws_k['C11'].alignment = align_center
    ws_k['C11'].number_format = '0.0%'
    
    # Add rank and status to the report card
    extra_k_rows = [
        (12, 'رتبه در استان:', '=IFERROR(VLOOKUP($C$2, \'داشبورد مانیتورینگ نواحی\'!$B$31:$U$62, 20, FALSE), "-")', '@'),
        (13, 'سطح ارزیابی عملکرد:', '=IFERROR(VLOOKUP($C$2, \'داشبورد مانیتورینگ نواحی\'!$B$31:$V$62, 21, FALSE), "-")', '@'),
        (14, 'تعداد حوزه مقاومت:', '=IFERROR(VLOOKUP($C$2, \'پایگاه داده حد انتظار\'!$A$2:$B$33, 2, FALSE), 0)', '#,##0'),
        (15, 'حدانتظار اعضای انجمن:', '=IFERROR(VLOOKUP($C$2, \'پایگاه داده حد انتظار\'!$A$2:$H$33, 8, FALSE), 0)', '#,##0')
    ]
    for r_k, lbl, fmla, n_fmt in extra_k_rows:
        ws_k.row_dimensions[r_k].height = 22
        cell_lbl = ws_k.cell(row=r_k, column=2)
        cell_lbl.value = lbl
        cell_lbl.font = Font(name=FONT_FAMILY, size=10, bold=True, color='181C1F')
        cell_lbl.fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        cell_lbl.alignment = align_right
        cell_lbl.border = border_thin_gray
        
        cell_val = ws_k.cell(row=r_k, column=3)
        cell_val.value = fmla
        cell_val.font = Font(name=FONT_FAMILY, size=10, bold=True, color='1B365D')
        cell_val.alignment = align_center
        cell_val.border = border_thin_gray
        cell_val.number_format = n_fmt

    # Set rightToLeft for all remaining sheets
    for s_name in ['گزارش عملکرد ماهانه', 'پایگاه داده حد انتظار']:
        if s_name in wb.sheetnames:
            wb[s_name].sheet_view.rightToLeft = True

    # Reorder sheets logically
    wb._sheets = [
        wb['داشبورد مانیتورینگ نواحی'],
        wb['کارنامه هوشمند'],
        wb['گزارش عملکرد ماهانه'],
        wb['پایگاه داده حد انتظار']
    ]

    # Save workbook
    wb.save(filename)
    print("Dashboard created and workbook saved successfully!")

if __name__ == '__main__':
    build()
