#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""کارکرد ماهانه — تولید خودکار فرم اکسل شرح فعالیت روزانه.

اجرا به دو صورت:
    ۱) تعاملی:      python karkard.py
    ۲) آرگومانی:    python karkard.py --month 5 --year 1405 --name "محسن ابوطالبیان"

پیش‌نیاز:  pip install -r requirements.txt
"""

from __future__ import annotations

import argparse
import json
import math
import random
import sys
from pathlib import Path

import jdatetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"

MONTH_NAMES = [
    "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
    "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند",
]
# jdatetime.weekday(): شنبه=0 ... جمعه=6
WEEKDAY_NAMES = ["شنبه", "یکشنبه", "دوشنبه", "سه‌شنبه", "چهارشنبه", "پنجشنبه", "جمعه"]

FULL_DAY_DUTY = 7 * 60 + 33   # موظفی روزانه: ۷ ساعت و ۳۳ دقیقه = ۴۵۳ دقیقه
THU_HALF_DUTY = 4 * 60        # موظفی پنجشنبه نیمه‌وقت: ۴ ساعت = ۲۴۰ دقیقه

HOLIDAY_WORK_MIN = 60         # کف دورکاری روز تعطیل: ۱ ساعت
HOLIDAY_WORK_MAX = 150        # سقف دورکاری روز تعطیل: ۲:۳۰
HOLIDAY_ENTRY_MIN = 540       # ۹:۰۰
HOLIDAY_ENTRY_MAX = 690       # ۱۱:۳۰

FA_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


# ---------------------------------------------------------------- ابزارها
def fa(value) -> str:
    """تبدیل ارقام لاتین به فارسی."""
    return str(value).translate(FA_DIGITS)


def fmt_clock(total_minutes: int) -> str:
    """قالب ساعت ورود/خروج: ۰۷:۳۵"""
    h, m = divmod(total_minutes, 60)
    return fa(f"{h:02d}:{m:02d}")


def fmt_duration(total_minutes: int) -> str:
    """قالب جمع ساعت: ۸:۳۰"""
    h, m = divmod(total_minutes, 60)
    return fa(f"{h}:{m:02d}")


def fmt_long(total_minutes: int) -> str:
    """قالب طولانی: ۲۱۸ ساعت و ۲۵ دقیقه"""
    h, m = divmod(total_minutes, 60)
    return f"{fa(h)} ساعت و {fa(m)} دقیقه"


def month_length(year: int, month: int) -> int:
    if month <= 6:
        return 31
    if month <= 11:
        return 30
    try:
        jdatetime.date(year, 12, 30)
        return 30
    except ValueError:
        return 29


def load_json(name: str):
    path = DATA_DIR / name
    if not path.exists():
        sys.exit(f"خطا: فایل {path} پیدا نشد.")
    return json.loads(path.read_text(encoding="utf-8"))


def parse_month(value: str) -> int:
    value = value.strip()
    if value in MONTH_NAMES:
        return MONTH_NAMES.index(value) + 1
    try:
        num = int(value.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")))
    except ValueError:
        sys.exit(f"خطا: ماه نامعتبر است: {value}")
    if not 1 <= num <= 12:
        sys.exit("خطا: شماره ماه باید بین ۱ تا ۱۲ باشد.")
    return num


# ---------------------------------------------------------------- موتور ساعات
def distribute_hours(work_days: list[dict], target: int, rng: random.Random) -> None:
    """توزیع ساعت ماهانه روی روزهای کاری به‌صورت تصادفی و سپس تنظیم دقیق به هدف.

    هر روز کاری دیکشنری با کلید kind است: 'full' یا 'thu'.
    نتیجه در کلید minutes هر روز ذخیره می‌شود.
    """
    bounds = {"full": (465, 555, 450, 600), "thu": (210, 270, 180, 300)}
    for day in work_days:
        lo, hi, _, _ = bounds[day["kind"]]
        span = (hi - lo) // 5
        day["minutes"] = lo + rng.randrange(span + 1) * 5

    diff = target - sum(d["minutes"] for d in work_days)
    guard = 0
    while diff != 0 and guard < 20000:
        guard += 1
        step = 5 if diff > 0 else -5
        cands = [
            d for d in work_days
            if bounds[d["kind"]][2] <= d["minutes"] + step <= bounds[d["kind"]][3]
        ]
        if not cands:
            break
        pick = rng.choice(cands)
        pick["minutes"] += step
        diff -= step

    final = sum(d["minutes"] for d in work_days)
    if final != target:
        # اختلاف باقی‌مانده روی اولین روز ممکن سرشکن می‌شود
        rest = target - final
        for day in work_days:
            lo_abs, hi_abs = bounds[day["kind"]][2], bounds[day["kind"]][3]
            if lo_abs <= day["minutes"] + rest <= hi_abs:
                day["minutes"] += rest
                break


# ---------------------------------------------------------------- موتور متن
def build_descriptions(days: list[dict], activities: dict, regions: list[str],
                        rng: random.Random) -> None:
    """ساخت شرح فعالیت هر روز. نتیجه در کلید desc ذخیره می‌شود."""
    routines = activities["routines"][:]
    rng.shuffle(routines)
    mains = (activities["rasad"] + activities["tolid"]
             + activities["enteshar"] + activities["team"])
    rng.shuffle(mains)
    starts = activities["start_month"][:]
    ends = activities["end_month"][:]
    rng.shuffle(starts)
    rng.shuffle(ends)
    hws = activities["holiday_work"][:]
    rng.shuffle(hws)
    region_pool = regions[:]
    rng.shuffle(region_pool)

    main_idx = start_idx = end_idx = region_idx = routine_idx = hw_idx = 0

    def next_region() -> str:
        nonlocal region_idx
        name = region_pool[region_idx % len(region_pool)]
        region_idx += 1
        return name

    def fill(text: str) -> str:
        while "{nahiye}" in text:
            text = text.replace("{nahiye}", next_region(), 1)
        return text

    work_days = [d for d in days if d["kind"] in ("full", "thu")]
    first3 = {d["day"] for d in work_days[:3]}
    last3 = {d["day"] for d in work_days[-3:]}

    for day in days:
        kind = day["kind"]
        if kind == "friday":
            day["desc"] = "جمعه (تعطیل هفتگی)"
            if day.get("holiday_title"):
                day["desc"] += f" ـ مصادف با تعطیل رسمی ({day['holiday_title']})"
        elif kind == "official":
            title = day["holiday_title"]
            day["desc"] = title if title.startswith("تعطیلی عمومی") \
                else f"تعطیل رسمی ({title})"
        elif kind == "thu_off":
            day["desc"] = "پنجشنبه (تعطیل هفتگی)"
        elif kind == "worked_holiday":
            if day["weekday"] == "جمعه":
                label = "جمعه (تعطیل هفتگی)"
                if day.get("holiday_title"):
                    label += f" ـ مصادف با تعطیل رسمی ({day['holiday_title']})"
            else:
                label = f"تعطیل رسمی ({day['holiday_title']})"
            n = 1 if rng.random() < 0.7 else 2
            items = []
            for _ in range(n):
                items.append(fill(hws[hw_idx % len(hws)]))
                hw_idx += 1
            day["desc"] = f"{label} ـ {'؛ '.join(items)} (دورکاری)"
        else:
            parts = [fill(routines[routine_idx % len(routines)])]
            routine_idx += 1
            if day["day"] in first3:
                parts.append(fill(starts[start_idx % len(starts)]))
                start_idx += 1
            if day["day"] in last3:
                parts.append(fill(ends[end_idx % len(ends)]))
                end_idx += 1
            need = 1 if kind == "thu" else 2
            for _ in range(need):
                parts.append(fill(mains[main_idx % len(mains)]))
                main_idx += 1
            day["desc"] = "؛ ".join(parts)


# ---------------------------------------------------------------- ساخت اکسل
PRIMARY = "1F4E5F"      # سرمه‌ای مایل به سبز (نوار عنوان)
HEADER_BLUE = "2E75B6"  # آبی سرستون‌ها
LIGHT_BAND = "D9E1F2"   # نوار روشن زیر عنوان
ZEBRA = "EAF1FA"        # سطرهای یکی‌درمیان
FILL_OFF = "E7E6E6"     # روزهای کاملاً تعطیل
FILL_HOLIDAY = "E2EFDA" # تعطیلات دارای دورکاری
FILL_TOTAL = "FFD966"   # سطر جمع
GRID = "B0B0B0"

THIN = Side(style="thin", color=GRID)
MED = Side(style="medium", color=PRIMARY)
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_range(ws, row: int, cols: str, font=None, fill=None, alignment=None, border=None):
    """اعمال استایل به تک‌تک سلول‌های یک محدوده (برای مربع‌های ادغام‌شده)."""
    from openpyxl.utils import column_index_from_string
    start, end = cols.split(":")
    for col in range(column_index_from_string(start), column_index_from_string(end) + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = border or BORDER_ALL
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if alignment:
            cell.alignment = alignment


def estimate_row_height(text: str) -> float:
    lines = max(1, math.ceil(len(text) / 90))
    return max(28.0, lines * 16.0)


def build_workbook(days: list[dict], year: int, month: int, person_title: str,
                   person_name: str, total_minutes: int, font_name: str,
                   show_duty: bool, duty_minutes: int) -> Workbook:
    month_name = MONTH_NAMES[month - 1]
    wb = Workbook()
    ws = wb.active
    ws.title = f"کارکرد {month_name}"
    ws.sheet_view.rightToLeft = True

    widths = {"A": 7, "B": 12, "C": 10, "D": 10, "E": 10, "F": 95}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    f_besmellah = Font(name=font_name, size=15, bold=True, color=PRIMARY)
    f_band = Font(name=font_name, size=14, bold=True, color="FFFFFF")
    f_sub = Font(name=font_name, size=12, color=PRIMARY)
    f_sig = Font(name=font_name, size=12, bold=True, color="FFFFFF")
    f_text = Font(name=font_name, size=12)
    f_head = Font(name=font_name, size=12, bold=True, color="FFFFFF")
    f_total = Font(name=font_name, size=13, bold=True, color=PRIMARY)
    f_off = Font(name=font_name, size=12, color="595959")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    desc_align = Alignment(horizontal="right", vertical="center", wrap_text=True)

    fill_primary = PatternFill("solid", fgColor=PRIMARY)
    fill_band = PatternFill("solid", fgColor=LIGHT_BAND)
    fill_head = PatternFill("solid", fgColor=HEADER_BLUE)
    fill_zebra = PatternFill("solid", fgColor=ZEBRA)
    fill_off = PatternFill("solid", fgColor=FILL_OFF)
    fill_holiday = PatternFill("solid", fgColor=FILL_HOLIDAY)
    fill_total = PatternFill("solid", fgColor=FILL_TOTAL)

    # --- سربرگ ---
    ws.merge_cells("A1:F1")
    ws["A1"] = "بسمه تعالی"
    style_range(ws, 1, "A:F", font=f_besmellah, alignment=center)
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:F2")
    ws["A2"] = "فرم ثبت شرح فعالیت روزانه نیروهای خرید خدمت مرکز فضای مجازی بسیج استان اصفهان"
    style_range(ws, 2, "A:F", font=f_band, fill=fill_primary, alignment=center)
    ws.row_dimensions[2].height = 30

    ws.merge_cells("A3:F3")
    ws["A3"] = (f"با احترام، بدینوسیله کارکرد {person_title} {person_name} "
                f"در {month_name} ماه {fa(year)} بر اساس جدول ذیل حضورتان ارسال می‌گردد:")
    style_range(ws, 3, "A:F", font=f_sub, fill=fill_band, alignment=center)
    ws.row_dimensions[3].height = 32

    style_range(ws, 4, "A:F", fill=fill_primary, alignment=center)
    ws.row_dimensions[4].height = 6

    # --- سرستون‌ها ---
    headers = ["ایام ماه", "روز هفته", "ساعت ورود", "ساعت خروج", "جمع ساعت",
               "ثبت شرح فعالیت دقیق روزانه (رصد و پایش / تولید محتوا / انتشار) "
               "با محوریت تیم‌سازی و توسعه کمی و کیفی"]
    for i, text in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=i, value=text)
        cell.font = f_head
        cell.fill = fill_head
        cell.alignment = center
        cell.border = BORDER_ALL
    ws.row_dimensions[5].height = 50

    # --- روزهای ماه ---
    row = 6
    zebra_on = False
    for day in days:
        kind = day["kind"]
        if kind in ("friday", "official", "thu_off"):
            fill, desc_font = fill_off, f_off
        elif kind == "worked_holiday":
            fill, desc_font = fill_holiday, f_text
        else:
            zebra_on = not zebra_on
            fill, desc_font = (fill_zebra if zebra_on else None), f_text
        off = kind in ("friday", "official", "thu_off")
        values = [
            fa(day["day"]),
            day["weekday"],
            "" if off else fmt_clock(day["entry"]),
            "" if off else fmt_clock(day["entry"] + day["minutes"]),
            "" if off else fmt_duration(day["minutes"]),
            day["desc"],
        ]
        for i, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=value)
            cell.font = desc_font if i == 6 else f_text
            cell.border = BORDER_ALL
            if fill:
                cell.fill = fill
            cell.alignment = desc_align if i == 6 else center
        ws.row_dimensions[row].height = estimate_row_height(day["desc"])
        row += 1
    last_day_row = row - 1

    # قاب متوسط دور جدول
    for c in range(1, 7):
        top = ws.cell(row=5, column=c)
        top.border = Border(top=MED, left=top.border.left, right=top.border.right,
                            bottom=top.border.bottom)
        bot = ws.cell(row=last_day_row, column=c)
        bot.border = Border(top=bot.border.top, left=bot.border.left, right=bot.border.right,
                            bottom=MED)
    for r in range(5, last_day_row + 1):
        left = ws.cell(row=r, column=1)
        left.border = Border(top=left.border.top, left=MED, right=left.border.right,
                             bottom=left.border.bottom)
        right = ws.cell(row=r, column=6)
        right.border = Border(top=right.border.top, left=right.border.left, right=MED,
                              bottom=right.border.bottom)

    # --- سطر جمع ---
    ws.merge_cells(f"A{row}:F{row}")
    ws.cell(row=row, column=1).value = f"جمع ساعات کارکرد در ماه: {fmt_long(total_minutes)}"
    style_range(ws, row, "A:F", font=f_total, fill=fill_total, alignment=center,
                border=Border(left=MED, right=MED, top=MED, bottom=MED))
    ws.row_dimensions[row].height = 30
    row += 1

    if show_duty:
        extra = total_minutes - duty_minutes
        ws.merge_cells(f"A{row}:F{row}")
        ws.cell(row=row, column=1).value = (
            f"موظفی ماه: {fmt_long(duty_minutes)} ـ اضافه‌کار: {fmt_long(extra)}")
        style_range(ws, row, "A:F", font=f_sub, alignment=center)
        ws.row_dimensions[row].height = 24
        row += 1

    row += 1  # یک سطر فاصله
    ws.row_dimensions[row - 1].height = 10

    # --- بلوک امضا ---
    sigs = [("A:B", "امضاء فرد"),
            ("C:D", "امضاء مسئول مرکز فضای مجازی"),
            ("E:F", "تأیید مرکز فضای مجازی بسیج استان")]
    for cols, text in sigs:
        a, b = cols.split(":")
        col0 = {"A": 1, "C": 3, "E": 5}[a]
        ws.merge_cells(f"{a}{row}:{b}{row}")
        ws.cell(row=row, column=col0).value = text
        style_range(ws, row, cols, font=f_sig, fill=fill_primary, alignment=center)
    ws.row_dimensions[row].height = 26
    row += 1
    for cols, _ in sigs:
        a, b = cols.split(":")
        ws.merge_cells(f"{a}{row}:{b}{row + 1}")
        style_range(ws, row, cols, alignment=center)
        style_range(ws, row + 1, cols, alignment=center)
    ws.row_dimensions[row].height = 38
    ws.row_dimensions[row + 1].height = 38
    last_row = row + 1

    # --- تنظیمات چاپ A4 ---
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "5:5"
    ws.print_area = f"A1:F{last_row}"
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.print_options.horizontalCentered = True
    footer_text = f"کارکرد {month_name} {fa(year)} ـ صفحه &P از &N"
    ws.oddFooter.center.text = footer_text
    ws.oddFooter.center.size = 9
    ws.evenFooter.center.text = footer_text
    ws.evenFooter.center.size = 9
    ws.freeze_panes = "A6"

    return wb


# ---------------------------------------------------------------- منطق اصلی
def generate(year: int, month: int, person_title: str, person_name: str,
             seed, factor_min: float, factor_max: float, fixed_factor,
             thursday_mode: str, font_name: str, show_duty: bool,
             output: str | None) -> Path:
    holidays_data = load_json("holidays.json")
    activities = load_json("activities.json")
    regions = load_json("regions.json")["regions"]

    year_key = str(year)
    if year_key not in holidays_data:
        sys.exit(f"خطا: تعطیلات سال {fa(year)} در پایگاه داده نیست. "
                 "ابتدا با tools/update_holidays.py آن را اضافه کنید.")
    if year in holidays_data.get("_meta", {}).get("estimated_years", []):
        print(f"⚠ هشدار: تعطیلات قمری سال {fa(year)} تخمینی است؛ با تقویم رسمی کنترل شود.")

    holidays = holidays_data[year_key]
    n_days = month_length(year, month)

    if seed is None:
        seed = random.randrange(1_000_000)
    print(f"بذر تصادفی (seed): {seed}")
    rng_hours = random.Random(f"{seed}-hours")
    rng_entry = random.Random(f"{seed}-entry")
    rng_text = random.Random(f"{seed}-text")
    rng_hol = random.Random(f"{seed}-holiday")

    # --- طبقه‌بندی روزها ---
    days: list[dict] = []
    for day in range(1, n_days + 1):
        weekday = jdatetime.date(year, month, day).weekday()  # شنبه=0
        title = holidays.get(f"{month}/{day}")
        if weekday == 6:
            kind = "friday"
        elif title is not None:
            kind = "official"
        elif weekday == 5:
            kind = {"half": "thu", "full": "full", "off": "thu_off"}[thursday_mode]
        else:
            kind = "full"
        days.append({"day": day, "weekday": WEEKDAY_NAMES[weekday],
                     "kind": kind, "holiday_title": title})

    # --- دورکاری یک‌سوم تعطیلات ---
    off_days = [d for d in days if d["kind"] in ("friday", "official")]
    rng_hol.shuffle(off_days)
    n_hol = max(1, int(len(off_days) / 3 + 0.5))
    for d in off_days[:n_hol]:
        d["kind"] = "worked_holiday"
        span = (HOLIDAY_WORK_MAX - HOLIDAY_WORK_MIN) // 5
        d["minutes"] = HOLIDAY_WORK_MIN + rng_hol.randrange(span + 1) * 5
        estep = (HOLIDAY_ENTRY_MAX - HOLIDAY_ENTRY_MIN) // 5
        d["entry"] = HOLIDAY_ENTRY_MIN + rng_hol.randrange(estep + 1) * 5
    h_minutes = sum(d["minutes"] for d in days if d["kind"] == "worked_holiday")

    work_days = [d for d in days if d["kind"] in ("full", "thu")]

    # --- موظفی و هدف (ضریب روی جمع کل شامل دورکاری تعطیلات اعمال می‌شود) ---
    n_full = sum(1 for d in work_days if d["kind"] == "full")
    n_thu = sum(1 for d in work_days if d["kind"] == "thu")
    duty = n_full * FULL_DAY_DUTY + n_thu * THU_HALF_DUTY
    factor = fixed_factor if fixed_factor is not None else rng_hours.uniform(factor_min, factor_max)
    work_target = round((duty * factor - h_minutes) / 5) * 5

    # --- توزیع ساعات و ورود/خروج ---
    distribute_hours(work_days, work_target, rng_hours)
    for day in work_days:
        day["entry"] = 440 + rng_entry.randrange(11) * 5  # ۷:۲۰ تا ۸:۱۰

    build_descriptions(days, activities, regions, rng_text)

    total = sum(d.get("minutes", 0) for d in days if d["kind"] in ("full", "thu", "worked_holiday"))

    # --- خروجی اکسل ---
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    month_name = MONTH_NAMES[month - 1]
    out_path = Path(output) if output else OUTPUT_DIR / f"کارکرد-{month_name}-{year}.xlsx"
    wb = build_workbook(days, year, month, person_title, person_name,
                        total, font_name, show_duty, duty)
    wb.save(out_path)

    # --- گزارش کنسول ---
    print(f"ماه: {month_name} {fa(year)} ـ روزهای ماه: {fa(n_days)} ـ "
          f"روز کامل: {fa(n_full)} ـ پنجشنبه کاری: {fa(n_thu)} ـ "
          f"تعطیل: {fa(len(off_days))} (دورکاری: {fa(n_hol)})")
    print(f"موظفی ماه: {fmt_long(duty)} ـ دورکاری تعطیلات: {fmt_long(h_minutes)}")
    print(f"ضریب اضافه‌کار: {fa(f'{factor:.2f}')} ـ جمع کارکرد: {fmt_long(total)} "
          f"({fa(f'{total / duty * 100:.1f}')}٪ موظفی)")
    print(f"✅ فایل ساخته شد: {out_path}")
    return out_path


def interactive() -> argparse.Namespace:
    print("─── کارکرد ماهانه (حالت تعاملی) ───")
    month_in = input("ماه (عدد ۱ تا ۱۲ یا نام ماه) [۵]: ").strip() or "5"
    year_in = input("سال [1405]: ").strip() or "1405"
    name = input("نام و نام خانوادگی [محسن ابوطالبیان]: ").strip() or "محسن ابوطالبیان"
    title = input("عنوان (برادر/خواهر) [برادر]: ").strip() or "برادر"
    seed_in = input("بذر تصادفی - خالی برای تصادفی (Enter): ").strip() or None
    ns = argparse.Namespace(
        month=month_in, year=int(year_in.translate(
            str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789"))),
        name=name, title=title, seed=seed_in,
        factor_min=1.10, factor_max=1.20, factor=None,
        thursday_mode="half", font="B Nazanin", show_duty=False, output=None,
    )
    return ns


def main(argv=None):
    parser = argparse.ArgumentParser(description="تولید فرم اکسل کارکرد ماهانه")
    parser.add_argument("--month", help="شماره یا نام ماه (مثل 5 یا مرداد)")
    parser.add_argument("--year", type=int, default=1405, help="سال (پیش‌فرض: 1405)")
    parser.add_argument("--name", default="محسن ابوطالبیان", help="نام و نام خانوادگی")
    parser.add_argument("--title", default="برادر", choices=["برادر", "خواهر"])
    parser.add_argument("--seed", default=None, help="بذر تصادفی برای بازتولید دقیق")
    parser.add_argument("--factor-min", type=float, default=1.10)
    parser.add_argument("--factor-max", type=float, default=1.20)
    parser.add_argument("--factor", type=float, default=None, help="ضریب ثابت اضافه‌کار")
    parser.add_argument("--thursday-mode", default="half", choices=["half", "full", "off"],
                        help="وضعیت پنجشنبه‌ها: نیمه‌وقت/کامل/تعطیل")
    parser.add_argument("--font", default="B Nazanin", help="فونت اکسل")
    parser.add_argument("--show-duty", action="store_true", help="نمایش موظفی و اضافه‌کار در اکسل")
    parser.add_argument("--output", default=None, help="مسیر فایل خروجی")
    args = parser.parse_args(argv)

    if args.month is None:
        try:
            args = interactive()
        except EOFError:
            parser.error("ماه مشخص نشده است (--month)")

    generate(year=args.year, month=parse_month(str(args.month)), person_title=args.title,
             person_name=args.name, seed=args.seed, factor_min=args.factor_min,
             factor_max=args.factor_max, fixed_factor=args.factor,
             thursday_mode=args.thursday_mode, font_name=args.font,
             show_duty=args.show_duty, output=args.output)


if __name__ == "__main__":
    main()
