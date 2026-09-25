# -*- coding: utf-8 -*-
"""
آزمون‌های خودکار گزارش‌یار (بدون نیاز به اینترنت یا مدل واقعی).

اجرا:  python3 test_pipeline.py
"""
import glob
import json
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import ai_engine            # noqa: E402
import gozaresh as gz       # noqa: E402

OUT = os.path.join(tempfile.gettempdir(), "gozaresh-test-out")   # پوشه‌ی موقت آزمون
DEMO = os.path.join(HERE, "demo-input")


def run_cli(*args) -> int:
    """اجرای برنامه با پارامترهای داده‌شده (مثل اجرای واقعی)."""
    return gz.main(list(args))


class TestDates(unittest.TestCase):
    def test_round_trip(self):
        for j in [(1405, 1, 1), (1405, 6, 31), (1403, 12, 29), (1400, 7, 15), (1404, 11, 30)]:
            gy, gm, gd = gz.jalali_to_gregorian(*j)
            self.assertEqual(gz.gregorian_to_jalali(gy, gm, gd), j, f"خطای تبدیل برای {j}")

    def test_range_parsing(self):
        r = gz.resolve_date_range("شهریور", "شهریور", 1405)
        self.assertIsNotNone(r["start"])
        self.assertIsNotNone(r["end"])
        self.assertLess(r["start"], r["end"])
        self.assertIn("شهریور", r["label"])
        r2 = gz.resolve_date_range("1405/06/01", "1405/06/31", 1405)
        self.assertEqual(r["start"], r2["start"])
        self.assertEqual(r["end"], r2["end"])

    def test_range_filter(self):
        posts = [{"id": 1, "ts": gz.jalali_to_unix(1405, 5, 20), "text": "x"},
                 {"id": 2, "ts": gz.jalali_to_unix(1405, 6, 10), "text": "y"},
                 {"id": 3, "ts": gz.jalali_to_unix(1405, 7, 2), "text": "z"}]
        r = gz.resolve_date_range("شهریور", "شهریور", 1405)
        got = [p["id"] for p in gz.filter_posts(posts, r["start"], r["end"])]
        self.assertEqual(got, [2])


class TestExtraction(unittest.TestCase):
    """استخراج داده از متن پست‌های فارسی (قاعده‌محور)."""

    def _rec(self, text):
        post = {"id": 5, "ts": gz.jalali_to_unix(1405, 6, 12), "text": gz.normalize_text(text),
                "channel": "test", "username": "test", "source_file": "t.json"}
        return gz.rule_record_from_post(post, 1)

    def test_full_post(self):
        rec = self._rec("برگزاری کارگاه سواد رسانه در ناحیه کاشان با حضور 45 نفر از "
                        "دانش‌آموزان و تدریس حمید گودرزی در دبیرستان شهید بهشتی")
        self.assertEqual(rec.county, "کاشان")
        self.assertEqual(rec.people, 45)
        self.assertEqual(rec.teacher, "حمید گودرزی")
        self.assertIn("دبیرستان", rec.place)
        self.assertEqual(rec.sheet, "حضوری")

    def test_persian_digits_and_live(self):
        rec = self._rec("پخش زنده‌ی کارگاه امنیت سایبری ناحیه نجف آباد با ۱۲۰ نفر بازدید؛ "
                        "تدریس مریم رضایی")
        self.assertEqual(rec.county, "نجف آباد")
        self.assertEqual(rec.people, 120)
        self.assertEqual(rec.sheet, "مجازی")

    def test_production(self):
        rec = self._rec("تولید کلیپ جنگ شناختی توسط ناحیه گلپایگان منتشر شد")
        self.assertEqual(rec.sheet, "تولیدات")
        self.assertEqual(rec.county, "گلپایگان")
        self.assertTrue(rec.production_kind)

    def test_no_county(self):
        rec = self._rec("یادآوری: مستندات در سامانه بارگذاری شود")
        self.assertEqual(rec.county, "")
        self.assertEqual(gz.county_label(rec), "نامشخص")


class TestAIGuards(unittest.TestCase):
    """کنترل‌های ایمنی هوش مصنوعی: نپذیرفتن خروجی نامعتبر."""

    def test_json_extraction(self):
        self.assertEqual(ai_engine.extract_json('```json\n[{"a":1}]\n```'), [{"a": 1}])
        self.assertEqual(ai_engine.extract_json('توضیح اضافه [{"b": 2},]'),
                         [{"b": 2}])
        self.assertIsNone(ai_engine.extract_json("متنی که جیسون نیست"))

    def test_polish_rejects_number_change(self):
        eng = ai_engine.AIEngine(backend="mock")

        class Bad(ai_engine.AIEngine):
            def complete(self, *a, **k):
                return ai_engine.AIResult(True, raw="## بخش\nتعداد اقدام‌ها ۹۹۹ عدد است.\n")

        bad = Bad(backend="mock")
        text = "## بخش\nتعداد اقدام‌ها ۱۲ عدد است.\n"
        result, warn = bad.polish_report(text, 100)
        self.assertEqual(result, text)
        self.assertIn("رد شد", warn)

    def test_polish_accepts_valid(self):
        class Good(ai_engine.AIEngine):
            def complete(self, *a, **k):
                return ai_engine.AIResult(
                    True, raw="## بخش\nدر این بازه ۱۲ اقدام انجام شد.\n")

        good = Good(backend="mock")
        text = "## بخش\nتعداد اقدام‌ها ۱۲ عدد است.\n"
        result, warn = good.polish_report(text, 100)
        self.assertEqual(warn, "")
        self.assertIn("۱۲", result)

    def test_polish_rejects_structure_collapse(self):
        """مدلی که سطرها/سرتیترها را به‌هم بریزد نباید پذیرفته شود."""
        text = ("# گزارش\n\n## مقدمه‌ی آماری\nدر این بازه ۱۲ اقدام ثبت شد.\n\n"
                "## سیمای کلی\n• **کاشان** — کارگاه (تاریخ ۱۴۰۵/۰۶/۱۲، جایگاه)\n")

        class Collapse(ai_engine.AIEngine):
            def complete(self, *a, **k):
                return ai_engine.AIResult(True, raw=" ".join(text.split("\n")))

        collapsed = Collapse(backend="mock")
        result, warn = collapsed.polish_report(text, 200)
        self.assertEqual(result, text)
        self.assertNotEqual(warn, "")

        class DropBullets(ai_engine.AIEngine):
            def complete(self, *a, **k):
                return ai_engine.AIResult(
                    True, raw="# گزارش\n\n## مقدمه‌ی آماری\nدر این بازه ۱۲ اقدام ثبت شد.\n\n"
                              "## سیمای کلی\nخلاصه‌ی کلی بدون جزئیات.\n")

        dropped = DropBullets(backend="mock")
        result, warn = dropped.polish_report(text, 200)
        self.assertEqual(result, text)

    def test_confidence_gate(self):
        rec = gz.Record(county_raw="", county="", month_key="1405-06", sheet="حضوری",
                        people=0, uid=1)
        low = {"اطمینان": 5, "ناحیه": "کاشان", "تعداد": 40}
        self.assertFalse(gz.apply_ai_extraction(rec, low))
        self.assertEqual(rec.county, "")
        self.assertEqual(rec.people, 0)
        high = {"اطمینان": 80, "ناحیه": "کاشان", "تعداد": 40, "مدرس": "علی محمدی"}
        self.assertTrue(gz.apply_ai_extraction(rec, high))
        self.assertEqual(rec.county, "کاشان")
        self.assertEqual(rec.people, 40)

    def test_numbers_preserved_helper(self):
        self.assertEqual(ai_engine.numbers_of("۴۵ نفر و 12 پست"),
                         ["45", "12"])


class TestPipeline(unittest.TestCase):
    """اجرای کامل خط تولید در سه حالت: اکسل، تلگرام، اکسل+تلگرام (با مدل آزمایشی)."""

    @classmethod
    def setUpClass(cls):
        if os.path.isdir(OUT):
            shutil.rmtree(OUT, ignore_errors=True)
        # ساخت داده‌ی نمونه در صورت نبودن (اسکریپت‌های نمونه‌ساز فقط یک‌بار لازم است)
        import subprocess
        if not glob.glob(os.path.join(DEMO, "*.xlsx")):
            subprocess.run([sys.executable, os.path.join(HERE, "make_demo_data.py")], check=True)
        if not os.path.exists(os.path.join(DEMO, "channel", "result.json")):
            subprocess.run([sys.executable, os.path.join(HERE, "make_demo_channel.py")], check=True)

    def _outputs(self, suffix):
        return glob.glob(os.path.join(OUT, "*" + suffix))

    def test_01_excel_only(self):
        code = run_cli("--input", DEMO, "--out", OUT, "--max-words", "500")
        self.assertEqual(code, 0)
        self.assertTrue(self._outputs(".docx"))
        md = open(self._outputs(".md")[0], encoding="utf-8").read()
        self.assertIn("منبع داده:** فایل‌های اکسل", md.replace("**", "**"))
        self.assertLessEqual(len(md.split()), 500)

    def test_02_telegram_only(self):
        code = run_cli("--input", DEMO, "--source", "telegram",
                       "--telegram", os.path.join(DEMO, "channel"),
                       "--from", "شهریور", "--to", "شهریور",
                       "--out", OUT, "--max-words", "600")
        self.assertEqual(code, 0)
        md = open(sorted(self._outputs(".md"))[-1], encoding="utf-8").read()
        self.assertIn("خروجی JSON کانال تلگرام", md)
        self.assertIn("پست", md)
        self.assertLessEqual(len(md.split()), 600)

    def test_03_telegram_with_ai_mock(self):
        code = run_cli("--input", DEMO, "--source", "telegram",
                       "--telegram", os.path.join(DEMO, "channel"),
                       "--from", "شهریور", "--to", "شهریور",
                       "--ai", "mock", "--ai-tasks", "extract,classify,impact",
                       "--out", OUT, "--max-words", "600")
        self.assertEqual(code, 0)
        md = open(sorted(self._outputs(".md"))[-1], encoding="utf-8").read()
        self.assertIn("تحلیل هوشمند", md)

    def test_04_both_with_fill_gaps(self):
        code = run_cli("--input", DEMO, "--source", "both",
                       "--telegram", os.path.join(DEMO, "channel"), "--tg-fill-gaps",
                       "--from", "شهریور", "--to", "شهریور",
                       "--out", OUT, "--no-xlsx", "--max-words", "700")
        self.assertEqual(code, 0)
        md = open(sorted(self._outputs(".md"))[-1], encoding="utf-8").read()
        self.assertIn("فایل‌های اکسل نواحی + خروجی تلگرام", md)

    def test_05_excel_sheets(self):
        code = run_cli("--input", DEMO, "--source", "telegram",
                       "--telegram", os.path.join(DEMO, "channel"),
                       "--from", "شهریور", "--to", "شهریور",
                       "--out", OUT, "--no-docx", "--max-words", "600")
        self.assertEqual(code, 0)
        import openpyxl
        wb = openpyxl.load_workbook(sorted(self._outputs(".xlsx"))[-1])
        self.assertIn("اقدامات", wb.sheetnames)
        self.assertIn("پست‌های بدون داده", wb.sheetnames)
        headers = [c.value for c in wb["اقدامات"][1]]
        self.assertIn("منبع", headers)
        self.assertIn("شماره پست", headers)

    def test_06_telegram_paths(self):
        """مسیر پوشه‌ی کانال یا خودِ فایل JSON باید مستقیم شناسایی شود."""
        folder = os.path.join(DEMO, "channel")
        code = run_cli("--input", folder, "--from", "شهریور", "--to", "شهریور",
                       "--no-docx", "--no-xlsx", "--max-words", "300")
        self.assertEqual(code, 0)
        md = open(sorted(self._outputs(".md"))[-1], encoding="utf-8").read()
        self.assertIn("خروجی JSON کانال تلگرام", md)

        file_path = os.path.join(folder, "result.json")
        code = run_cli("--input", file_path, "--from", "شهریور", "--to", "شهریور",
                       "--no-docx", "--no-xlsx", "--max-words", "300")
        self.assertEqual(code, 0)

    def test_07_one_page_compact(self):
        """گزارش یک‌صفحه‌ای (سقف ۳۰۰ واژه): همه‌ی محورها با یک اقدام مشخص."""
        code = run_cli("--source", "telegram", "--telegram", os.path.join(DEMO, "channel"),
                       "--from", "مرداد", "--to", "شهریور", "--max-words", "300",
                       "--out", OUT, "--no-docx", "--no-xlsx")
        self.assertEqual(code, 0)
        md = open(sorted(self._outputs(".md"))[-1], encoding="utf-8").read()
        self.assertLessEqual(len(md.split()), 300)
        axes = ["بحران و مسئله ویژه", "تولیدات و پویش‌های جریان‌ساز",
                "هم‌افزایی و اقدامات مشترک", "برنامه‌ها و رویدادهای ویژه"]
        for ax in axes:
            self.assertIn(ax, md)
        # هر محور باید دست‌کم یک اقدام مشخص (نه فقط آماره) داشته باشد
        blocks = md.split("## بخش: ")[1:]
        for b in blocks:
            if b.startswith("سایر اقدامات"):
                continue
            self.assertIn("• **", b, "محور بدون اقدام مشخص: " + b.split("\n")[0])

    def test_08_request_text_sets_cap(self):
        """متن «درخواست مدیر» باید سقف واژه را (حتی با ارقام ترکیبی 3۰۰) اعمال کند."""
        req = ("یک گزارش حداکثر یک‌صفحه‌ای (3۰۰ کلمه) از اقدامات مهم طی دو ماه مرداد و شهریور "
               "ارسال کنند؛ از ذکر اقدامات روتین خودداری شود.")
        code = run_cli("--source", "telegram", "--telegram", os.path.join(DEMO, "channel"),
                       "--from", "مرداد", "--to", "شهریور", "--request", req,
                       "--out", OUT, "--no-docx", "--no-xlsx")
        self.assertEqual(code, 0)
        md = open(sorted(self._outputs(".md"))[-1], encoding="utf-8").read()
        self.assertLessEqual(len(md.split()), 300)

    def test_09_no_data_excel(self):
        """فایل اکسل بدون داده باید با پیام روشن رد شود (نه خطای برنامه)."""
        empty = os.path.join(HERE, "demo-input", "empty-test.xlsx")
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.title = "سوادرسانه حضوری"
        wb.save(empty)
        try:
            code = run_cli("--input", empty, "--out", OUT)
            self.assertEqual(code, 1)
        finally:
            os.remove(empty)


if __name__ == "__main__":
    unittest.main(verbosity=2)
