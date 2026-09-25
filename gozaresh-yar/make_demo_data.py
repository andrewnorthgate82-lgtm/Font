# -*- coding: utf-8 -*-
"""
ساخت داده‌ی آزمایشی (فرضی) برای تست موتور گزارش‌یار.
این فایل صرفاً برای آزمایش ابزار است و هیچ داده‌ی واقعی در آن وجود ندارد.
خروجی: چند فایل اکسل با همان قالب «گزارش ... ناحیه.xlsx»
"""
import os
import random
import openpyxl
from openpyxl.styles import Font, Alignment

random.seed(7)

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "demo-input")
os.makedirs(OUT, exist_ok=True)

TEMPLATE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                        "گزارش شهریور ماه 1405 ناحیه.xlsx")

SHEETS = ["سوادرسانه حضوری", "سوادرسانه مجازی (لایو)", "توانمندسازی گردان",
          "سواد رسانه خلاقانه", "تولیدات "]

HEADERS = {
    "سوادرسانه حضوری": ["ردیف", "تاریخ", "نام ناحیه", "موضوع", "نام سخنران / مدرس",
                        "مدت (دقیقه)", "تعداد نفرات", "مکان برگزاری", "لینک مستندات"],
    "سوادرسانه مجازی (لایو)": ["ردیف", "تاریخ", "نام ناحیه", "موضوع", "نام سخنران / مدرس",
                               "مدت (دقیقه)", "تعداد نفرات", "بستر برگزاری", "لینک مستندات"],
    "توانمندسازی گردان": ["ردیف", "تاریخ", "برگزار کننده", "موضوع", "نام سخنران / مدرس",
                          "مدت (دقیقه)", "تعداد نفرات", "مکان برگزاری", "لینک مستندات", "نوع کلاس"],
    "سواد رسانه خلاقانه": ["ردیف", "تاریخ", "ناحیه برگزار کننده", "موضوع",
                           "تعداد نفرات / تعداد بازدید ", "بستر برگزاری", "لینک مستندات"],
    "تولیدات ": ["ردیف", "تاریخ محتوای تولید شده", "تولید کننده", "موضوع",
                 "نام تولید کننده محتوا", "نوع تولید محتوا", "تعداد صفحات", "لینک مستندات"],
}

TOPICS = ["سوادرسانه", "جنگ شناختی", "آموزش هوش مصنوعی", "آموزش تولید محتوا",
          "امنیت سایبری", "سواد رسانه و فضای مجازی", "آسیب‌های شبکه‌های اجتماعی"]
INSTRUCTORS = ["حمید گودرزی", "مریم رضایی", "علی محمدی", "زهرا کریمی", "سعید نوری",
               "فاطمه احمدی", "محسن جعفری"]
PLACES = ["مسجد جامع", "دبیرستان شهید بهشتی", "حسینیه اعظم", "سالن اجتماعات ناحیه",
          "پایگاه مقاومت", "مدرسه ابتدایی شهید فهمیده"]
PLATFORMS = ["تلگرام", "ایتا", "روبیکا", "اینستاگرام", "شاد", "سامانه‌نسرا", "واتساپ"]
KINDS = ["ناحیه محور حضوری", "ناحیه محور مجازی", "حوزه محور حضوری", "حوزه محور مجازی"]
PRODUCTION_KINDS = ["صفحه word", "اسلاید", "کلیپ", "موشن گرافیک", "اینفوگرافیک", "پوستر"]
TOPICS_TEXT = {
    "سوادرسانه": "برگزاری جلسه سواد رسانه با محوریت تحلیل اخبار جعلی",
    "جنگ شناختی": "تشریح ابعاد جنگ شناختی و راه‌های مقابله با آن",
    "آموزش هوش مصنوعی": "کارگاه کاربرد هوش مصنوعی در تولید محتوا",
    "آموزش تولید محتوا": "آموزش اصول تولید محتوای رسانه‌ای و ویرایش",
    "امنیت سایبری": "نکات امنیت سایبری و حفاظت از حساب‌های کاربری",
    "سواد رسانه و فضای مجازی": "آموزش سواد رسانه و مدیریت مصرف فضای مجازی",
    "آسیب‌های شبکه‌های اجتماعی": "بررسی آسیب‌های شبکه‌های اجتماعی در خانواده",
}


def build_file(county, rows_per_sheet, out_name):
    """rows_per_sheet: dict sheet_name -> list of dicts"""
    wb = openpyxl.load_workbook(TEMPLATE)
    for sh in SHEETS:
        ws = wb[sh]
        rows = rows_per_sheet.get(sh, [])
        for i, row in enumerate(rows, start=1):
            r = i + 1
            ws.cell(r, 1).value = i
            if sh == "سوادرسانه حضوری":
                vals = [row["date"], county, row["topic"], row["teacher"], row["mins"],
                        row["people"], row["place"], row["link"]]
            elif sh == "سوادرسانه مجازی (لایو)":
                vals = [row["date"], county, row["topic"], row["teacher"], row["mins"],
                        row["people"], row["platform"], row["link"]]
            elif sh == "توانمندسازی گردان":
                vals = [row["date"], row["org"], row["topic"], row["teacher"], row["mins"],
                        row["people"], row["place"], row["link"], row["kind"]]
            elif sh == "سواد رسانه خلاقانه":
                vals = [row["date"], county, row["topic"], row["people"],
                        row["platform"], row["link"]]
            else:  # تولیدات
                vals = [row["date"], row["org"], row["topic"], row["producer"],
                        row["kind"], row["pages"], row["link"]]
            for j, v in enumerate(vals, start=2):
                c = ws.cell(r, j)
                c.value = v
                c.font = Font(name="b titr", size=12)
                c.alignment = Alignment(horizontal="center", vertical="center")
    path = os.path.join(OUT, out_name)
    wb.save(path)
    return path


def gen_rows(n, sh, county, tag):
    rows = []
    for i in range(n):
        d = f"1405/06/{random.randint(1, 31):02d}"
        topic = random.choice(TOPICS)
        link = f"https://t.me/soad_resaneh_isfahan/{random.randint(100, 999)}"
        if sh in ("سوادرسانه حضوری", "توانمندسازی گردان"):
            rows.append(dict(date=d, topic=topic, teacher=random.choice(INSTRUCTORS),
                             mins=random.choice([45, 60, 90, 120]),
                             people=random.randint(15, 120),
                             place=random.choice(PLACES), link=link,
                             org=random.choice(["ناحیه", "حوزه", "استان"]),
                             kind=random.choice(KINDS)))
        elif sh == "سوادرسانه مجازی (لایو)":
            rows.append(dict(date=d, topic=topic, teacher=random.choice(INSTRUCTORS),
                             mins=random.choice([30, 45, 60]),
                             people=random.randint(50, 900),
                             platform=random.choice(PLATFORMS), link=link))
        elif sh == "سواد رسانه خلاقانه":
            rows.append(dict(date=d, topic=topic, people=random.randint(100, 5000),
                             platform=random.choice(PLATFORMS), link=link))
        else:
            rows.append(dict(date=d, topic=topic, org=random.choice(["ناحیه", "ستاد استان"]),
                             producer=random.choice(INSTRUCTORS),
                             kind=random.choice(PRODUCTION_KINDS),
                             pages=random.randint(1, 4), link=link))
    return rows


if __name__ == "__main__":
    # سه ناحیه با سطح مشارکت متفاوت (فرضی)
    plan = {
        "آران و بیدگل": {"سوادرسانه حضوری": 14, "سوادرسانه مجازی (لایو)": 9,
                          "توانمندسازی گردان": 11, "سواد رسانه خلاقانه": 6, "تولیدات ": 5},
        "کاشان": {"سوادرسانه حضوری": 22, "سوادرسانه مجازی (لایو)": 13,
                   "توانمندسازی گردان": 16, "سواد رسانه خلاقانه": 8, "تولیدات ": 7},
        "شهرضا": {"سوادرسانه حضوری": 8, "سوادرسانه مجازی (لایو)": 4,
                   "توانمندسازی گردان": 5, "سواد رسانه خلاقانه": 3, "تولیدات ": 2},
    }
    for county, counts in plan.items():
        rows = {sh: gen_rows(n, sh, county, county) for sh, n in counts.items()}
        p = build_file(county, rows, f"{county}.xlsx")
        print("ساخته شد:", os.path.basename(p))
